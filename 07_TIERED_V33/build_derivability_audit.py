#!/usr/bin/env python3
"""
Build PRIMARY_OUTCOME_DERIVABILITY_AUDIT_v33 from the v32 master workbook.

Clinical estimand (defined independently of reporting format):
  Cumulative postoperative opioid consumption during the first 24 hours after
  surgery, standardized to intravenous morphine milligram equivalents (IV MME)
  where a defensible prespecified conversion is possible.

Tiers are assigned SEPARATELY for two metrics, because they have different
requirements:
  * tier_mme  -- the absolute IV MME estimand (needs absolute mass + a sourced
                 conversion factor)
  * smd_eligible -- the standardized (Hedges' g) metric, which is scale
                 invariant and therefore tolerates weight-normalized units
                 (mg/kg, ug/kg) that are NOT usable for absolute MME.

v32 is read-only. Nothing here writes back to it.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
V32 = ROOT / "TEAS EA Verification" / "TEAS_EA_RECONCILED_MASTER_DATA_v32_FINAL_LOCK_READY.xlsx"
OUT = ROOT / "07_TIERED_V33"

# Sourced IV MME conversion factors already verified in this project
# (BC Ministry of Health palliative equianalgesic table; FDA/Pfizer sufentanil
# label). Anything not in this dict has NO sourced factor and cannot be
# converted to absolute IV MME.
MME_FACTORS = {
    "mg morphine": 1.0, "mg iv morphine": 1.0, "mg mme": 1.0,
    "mg hydromorphone": 5.0,
    "ug sufentanil": 1.0, "µg sufentanil": 1.0,
    "mg fentanyl": 100.0,
}
NO_SOURCED_FACTOR = ("tramadol", "pethidine", "meperidine", "butorphanol",
                     "nalbuphine", "oxycodone", "dezocine")

COMPARATOR_RULES = [
    # (needle, category) -- order matters; first match wins. Case-insensitive.
    #
    # Inert-sham markers are tested FIRST. A device name such as "TENS" or an
    # electrode label does not make a comparator "active": Chen 1998's control
    # is "Sham ST36 TENS (0 mA)", i.e. the TENS device is applied at zero
    # current, which is an inert sham. Matching the device name before the
    # 0 mA / sham marker would misfile a genuine sham-controlled trial as an
    # active electrical comparator and corrupt the comparator strata.
    ("0 ma", "1_sham_placebo"),
    ("zero-current", "1_sham_placebo"),
    ("zero current", "1_sham_placebo"),
    ("no-current", "1_sham_placebo"),
    ("no current", "1_sham_placebo"),
    ("no stimulation", "1_sham_placebo"),
    ("nonpenetrating", "1_sham_placebo"),
    ("sub-sensory sham", "1_sham_placebo"),
    ("placebo", "1_sham_placebo"),
    # Active electrical comparators: real current delivered at a control site.
    ("nonacupoint aes", "2_active_electrical"),
    ("nonmeridian aes", "2_active_electrical"),
    ("active nonacupoint", "2_active_electrical"),
    ("incision-periphery tens", "2_active_electrical"),
    ("nonacupoint", "2_active_electrical"),
    ("non-acupoint", "2_active_electrical"),
    ("nonmeridian", "2_active_electrical"),
    ("sham", "1_sham_placebo"),
    ("usual care", "3_usual_care"),
    ("standard care", "3_usual_care"),
    ("no-aes", "3_usual_care"),
    ("no aes", "3_usual_care"),
    ("pca alone", "3_usual_care"),
    ("pca only", "3_usual_care"),
    ("eras", "3_usual_care"),
    ("no acupuncture", "3_usual_care"),
    ("no electrical-stimulation intervention", "3_usual_care"),
    ("control", "3_usual_care"),
]


def comparator_category(text: str) -> str:
    t = (text or "").lower()
    for needle, cat in COMPARATOR_RULES:
        if needle in t:
            return cat
    return "9_unclassified"


def mme_factor_for(unit: str):
    u = (unit or "").strip().lower()
    for drug in NO_SOURCED_FACTOR:
        if drug in u:
            return None, f"no sourced IV MME factor located for {drug}"
    if "/kg" in u:
        return None, "weight-normalized unit; absolute dose needs individual body weights"
    if "ml" in u and "solution" in u:
        return None, "solution volume; drug concentration not reported"
    for key, val in MME_FACTORS.items():
        if u == key or u.startswith(key):
            return val, "BC Ministry of Health equianalgesic table / FDA sufentanil label"
    return None, "unit not mapped to a sourced conversion factor"


# ---------------------------------------------------------------------------
# Manual, PDF-verified overrides. Every entry here was checked against the
# source PDF in this pass; `verified` records the exact locator.
# ---------------------------------------------------------------------------
OVERRIDES = {
    "GAO22_TEAS_vs_SHAM_SUFPUMP24": dict(
        tier_mme="C", smd_eligible="No",
        transformation="None. Exact 0-24 h cumulative pump sufentanil, median (IQR).",
        assumptions="None for the value itself. Mean/SD transformation (Wan/Luo) is NOT "
                    "defensible here: the 25th percentile is exactly 0.0 in BOTH arms, so the "
                    "distribution is zero-inflated and strongly non-normal, violating the "
                    "normality assumption those estimators rely on.",
        reason="Exact clinical estimand IS reported. Reporting format (median/IQR) is "
               "incompatible with the mean-difference model, not the estimand. Enters S2 via a "
               "distribution-compatible method; must NOT be transformed to mean/SD.",
        verified="pdf-2.pdf Table: '24 h analgesic pump, ug 33.0 (0.0, 50.0) vs 30.0 (0.0, 60.0) P=0.197'",
    ),
    "CHEN15Q_TEAS_vs_CTRL_RESCUEMORPH24": dict(
        tier_mme="B+C", smd_eligible="No",
        transformation="Deterministic: rescue count x fixed 2 mg IV morphine. Multiplication by a "
                       "positive constant is order-preserving, so median and IQR transform exactly "
                       "(median 1 (1-3) -> 2 (2-6); 3.5 (2-7) -> 7 (4-14) mg IV morphine).",
        assumptions="None. Fixed unit dose stated in source; morphine is the IV MME reference "
                    "standard (factor 1.0); rescue morphine was the only postoperative opioid.",
        reason="Deterministically derivable to exact mg IV MME (Tier B), but the resulting "
               "distribution is median/IQR (Tier C). Enters S1 only if the S1 method accepts "
               "quantile data; otherwise S2.",
        verified="037_chen_2015_thyroidectomy_lund.pdf: 'intravenous morphine 2 mg as rescue "
                 "analgesia to maintain a VAS score < 4'; Table: 'Cumulative number of rescue "
                 "analgesia 1 (1-3) vs 3.5 (2-7) P=0.004'",
    ),
    "CHEN15H_TEAS_vs_SHAM_SUFKG24": dict(
        tier_mme="E", smd_eligible="Yes (ug/kg, median/IQR - needs distribution-compatible method)",
        transformation="Deterministic in ug/kg: bolus count x fixed 0.05 ug/kg. Absolute ug "
                       "requires individual body weights, which are not reported.",
        assumptions="Absolute IV MME would require an assumed or group-average body weight - "
                    "explicitly prohibited.",
        reason="Valid weight-normalized endpoint; NOT convertible to absolute IV MME. Report "
               "separately in ug/kg; excluded from the absolute-MME primary and from S0-S3.",
        verified="040_chen_2015_hyperalgesia_lund.pdf: 'sufentanil 0.05 ug/kg as a bolus (VNRS "
                 "score >= 4)'; T5 = 24 h after surgery",
    ),
    "GU19_OPIOID24": dict(
        tier_mme="D", smd_eligible="Pending digitization",
        transformation="PCIA solution mL -> ug sufentanil is exact and deterministic: 100 ug "
                       "sufentanil in 100 mL = 1.0 ug/mL, so mL and ug are numerically equal. The "
                       "blocker is that the 24 h value itself is only plotted in Fig. 4.",
        assumptions="None for the unit conversion. Digitization must be formal and dual-reviewed; "
                    "values must not be visually estimated.",
        reason="Exact 0-24 h endpoint exists but only graphically. Text reports 4 h, 8 h and 36 h "
               "numerically and skips 24 h. NOTE: PCIA included a fixed 2 mL/h basal infusion, so a "
               "large protocol-mandated component is common to both arms and dilutes between-group "
               "contrast; record this with any digitized value.",
        verified="covidence_1471_full_article.pdf: 'PCIA devices with 100 ug sufentanil ... in 100 mL "
                 "(bolus 0.5 mL, lock-out 15 min, basal infusion 2 ml/h)'; section 3.3 reports 4 h, "
                 "8 h, 36 h only; 'Fig. 4 Comparison of analgesic consumption within 36 h'",
    ),
    "SONG20_TEAS_vs_SHAM_BUTORPCA24_DERIVED": dict(
        tier_mme="E", smd_eligible="No",
        transformation="Reported outcome is 'total number of PCA pump presses'. Converting to drug "
                       "mass requires presses == delivered boluses.",
        assumptions="TWO independent blockers. (1) Device semantics: the source records PRESSES and "
                    "documents a 15-min lockout, so demands and deliveries can differ; the source "
                    "never reports a delivered dose. (2) Even with resolved semantics, butorphanol "
                    "has no sourced IV MME equivalence in this project's reference tables.",
        reason="Assumption-dependent on both counts; excluded from the strict quantitative primary "
               "synthesis. Comparator is also an active electrical control, not sham.",
        verified="getfile.php-6.pdf: '4 mg butorphanol and 2 g propacetamol in 100 mL saline, every "
                 "pump press resulting in a 2 mL infusion, with a 15-min lockout interval'; "
                 "'Record the total number of PCA pump presses ... within 24 hours'; Table 3 "
                 "30.12+/-4.3 vs 36.91+/-2.4",
    ),
    "OZT19_TAES_vs_CTRL_TRAM24": dict(
        tier_mme="E", smd_eligible="Yes (tramadol-only endpoint, native mg)",
        transformation="None needed for mg tramadol. Absolute IV MME needs a tramadol conversion "
                       "factor.",
        assumptions="No sourced parenteral tramadol:morphine ratio exists in the reference family "
                    "used by this project (BC Ministry of Health and BC Renal tables omit tramadol "
                    "and pethidine entirely; verified 2026-09-07). Separately, tramadol alone is not "
                    "complete opioid exposure because IM rescue pethidine was also given.",
        reason="Exact 0-24 h mean/SD in native units, but not convertible to IV MME with a sourced "
               "factor, and not complete opioid exposure. Overall RoB High.",
        verified="covidence_505_full_article.pdf: 'Tramadol HCl (0-24 h) X+/-SD 257.68+/-65.55 | "
                 "228.40+/-87.89 | 357.81+/-123.70 .001'; 'Pethidine HCl (0-24 h) 33.75+/-56.55 | "
                 "20.66+/-27.89 | 33.12+/-42.69 .785'",
    ),
    "OZT19_TAES_vs_CTRL_TOTALOP24": dict(
        tier_mme="E", smd_eligible="No",
        transformation="Combined tramadol + pethidine total. Mean is additive after conversion; the "
                       "SD is not.",
        assumptions="SD(total) = sqrt(SD_t^2 + SD_p^2) requires zero within-person covariance, which "
                    "is unknown and cannot be assumed. Raw participant data unavailable.",
        reason="Combined-total sampling variance is not recoverable. Excluded from all quantitative "
               "tiers.",
        verified="covidence_505_full_article.pdf, same table as above (components reported separately)",
    ),
    "COURA11_EA_vs_SHAM": dict(
        tier_mme="E", smd_eligible="Yes (ug/kg, mean/SD)",
        transformation="ug/kg -> absolute ug would require body weight.",
        assumptions="Body weights not reported; a 70 kg assumption is explicitly prohibited. "
                    "Additionally the source permitted 'supplementary doses of morphine or fentanyl', "
                    "so reported fentanyl is not complete opioid exposure.",
        reason="Weight-normalized and incomplete-capture; excluded from the absolute-MME primary. "
               "Legitimate on the scale-invariant SMD metric, with the incomplete-capture caveat.",
        verified="covidence_819_full_article.pdf: 'Table 2 Fentanyl dosage during first 24 h after "
                 "operation'; 13.1+/-2.2 vs 16.3+/-1.6 ug/kg; 'supplementary doses of morphine or "
                 "fentanyl were allowed'",
    ),
    "ZHANG25_TEAS_vs_SHAM_SUFPOD1": dict(
        tier_mme="E", smd_eligible="No (window mismatch)",
        transformation="None available. The reported window is not the estimand's window.",
        assumptions="Treating POD 1 as the cumulative first 24 h is unsupported: the source defines "
                    "'postoperative days (PODs) 0 (day of surgery), 1, and 2', so POD 1 is a "
                    "separate calendar day that excludes day-of-surgery consumption and extends "
                    "past 24 h post-incision by an unknown amount.",
        reason="Window mismatch, not a reporting-format problem. Must be removed from ALL 0-24 h "
               "analyses including the broader SMD set that currently contains it.",
        verified="109551.pdf: 'assessed on postoperative days (PODs) 0 (day of surgery), 1, and 2'; "
                 "'total sufentanil consumption ... on POD 1'; Table 50.53+/-4.46 vs 53.79+/-5.14",
    ),
    "NTRI14_EA_vs_SHAM_TRAMSUM": dict(
        tier_mme="E", smd_eligible="No",
        transformation="None available.",
        assumptions="Reported total tramadol includes scheduled/background tramadol, not only "
                    "demand-driven postoperative consumption; rescue morphine quantity is not "
                    "reported. Tramadol also has no sourced IV MME factor.",
        reason="Endpoint is not the estimand (protocol-mandated background dosing included) and is "
               "not convertible to IV MME.",
        verified="v32 Opioid_24h_Candidates; PDF not re-verified this pass",
    ),
    "JIN23_2HZ_vs_SHAM_VOL24": dict(
        tier_mme="E", smd_eligible="Yes (mL proxy; scale cancels in SMD) - multi-arm",
        transformation="mL of PCIA solution -> fentanyl mass requires the solution concentration, "
                       "which the main article does not report.",
        assumptions="Absolute MME would require assuming a concentration. NOTE: within this study "
                    "both arms drew from the same solution, so mL is linear in fentanyl mass with a "
                    "common unknown constant; that constant cancels in a standardized effect, so the "
                    "contrast is usable on the SMD metric even though absolute MME is not "
                    "recoverable.",
        reason="No sourced absolute conversion; shares a sham arm with the 20/100 Hz contrast.",
        verified="v32 Opioid_24h_Candidates; concentration absent from main article per v32 QC",
    ),
    "JIN23_20100_vs_SHAM_VOL24": dict(
        tier_mme="E", smd_eligible="Yes (mL proxy; scale cancels in SMD) - shared sham arm",
        transformation="As above.",
        assumptions="As above; additionally this contrast shares the sham arm with the 2 Hz contrast "
                    "and must not be counted as an independent study unit.",
        reason="No sourced absolute conversion; correlated multi-arm contrast.",
        verified="v32 Opioid_24h_Candidates; concentration absent from main article per v32 QC",
    ),
    "SIM02_PREOP_vs_PLACEBO_MORPH24": dict(
        tier_mme="E", smd_eligible="Yes (mg/kg, mean/SD)",
        transformation="mg/kg -> absolute mg would require body weight.",
        assumptions="Body weights not reported; group-average reconstruction prohibited.",
        reason="Weight-normalized; excluded from absolute-MME primary, legitimate on SMD. "
               "Shared placebo control with the postoperative-EA arm: must not be double-counted.",
        verified="v32 Opioid_24h_Candidates (mg/kg morphine, 0-24 h); PDF not re-verified this pass",
    ),
    "SIM02_POSTOP_vs_PLACEBO_MORPH24": dict(
        tier_mme="E", smd_eligible="Yes (mg/kg, mean/SD) - shared control",
        transformation="mg/kg -> absolute mg would require body weight.",
        assumptions="Body weights not reported. Shares the placebo arm with the preoperative-EA "
                    "contrast.",
        reason="Same as the preoperative arm; additionally a correlated multi-arm contrast.",
        verified="v32 Opioid_24h_Candidates (mg/kg morphine, 0-24 h); PDF not re-verified this pass",
    ),
}


def classify(row: dict) -> dict:
    cid = (row.get("Comparison ID") or "").strip()
    unit = row.get("Unit") or ""
    dtype = (row.get("Data type") or "").strip()
    exact = (row.get("Exact 0-24 h?") or "").strip()
    readiness = (row.get("Analysis readiness") or "").strip()
    qc = (row.get("Source-QC issue") or "").strip()

    factor, factor_note = mme_factor_for(unit)
    out = dict(
        tier_mme="", smd_eligible="", transformation="", assumptions="",
        reason="", verified="not individually PDF-verified in this pass",
        mme_factor=("" if factor is None else factor), mme_factor_source=factor_note,
    )

    if cid in OVERRIDES:
        out.update(OVERRIDES[cid])
        if out.get("mme_factor") == "" and factor is not None:
            out["mme_factor"] = factor
        return out

    # rule-based fallback
    if "DUPLICATE HOLD" in readiness.upper() or "overlapping" in qc.lower():
        out.update(tier_mme="F", smd_eligible="No",
                   reason="Probable overlapping publication/cohort; not an independent study unit "
                          "until linkage is resolved.")
        return out
    if exact not in ("Yes", "Yes — derived", "Conditional"):
        out.update(tier_mme="F", smd_eligible="No",
                   reason="Not an exact cumulative 0-24 h postoperative opioid endpoint.")
        return out
    if dtype in ("Not reported", "Not calculable", "Not applicable / not reported"):
        out.update(tier_mme="F", smd_eligible="No",
                   reason="Exact 0-24 h endpoint is not quantitatively reported.")
        return out
    if "Graph-only" in dtype:
        out.update(tier_mme="D", smd_eligible="Pending digitization",
                   reason="Exact endpoint present only in a figure; formal digitization required.")
        return out
    if dtype == "Median/IQR":
        out.update(tier_mme="C", smd_eligible="No",
                   reason="Exact endpoint reported in an alternative distributional form.")
        return out
    if dtype in ("Mean/SD", "Derived mean/SD"):
        if factor is None:
            out.update(tier_mme="E", smd_eligible="Yes" if "/kg" in unit.lower() else "Check",
                       reason=f"Mean/SD available but not convertible to absolute IV MME: {factor_note}.")
            return out
        out.update(tier_mme="A", smd_eligible="Yes",
                   transformation=f"unit x {factor} mg IV MME",
                   reason="Exact 0-24 h cumulative endpoint, mean/SD, sourced absolute MME conversion.")
        return out
    out.update(tier_mme="F", smd_eligible="No", reason=f"Unhandled data type: {dtype}")
    return out


def main() -> int:
    wb = openpyxl.load_workbook(V32, data_only=True)
    ws = wb["Opioid_24h_Candidates"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]

    # current pipeline membership, to record what actually moves
    prim = {}
    with (ROOT / "06_FINAL_ANALYSIS_V26" / "01_DATA" / "opioid_24h_primary.csv").open(encoding="utf-8") as f:
        for r in csv.DictReader(f):
            prim[r["comparison_id"]] = (r["inc_primary"], r["inc_sens"])

    audit = []
    for r in rows:
        cid = (r.get("Comparison ID") or "").strip()
        c = classify(r)
        v32_prim, v32_sens = prim.get(cid, ("", ""))
        audit.append({
            "study": r.get("Study"),
            "comparison_id": cid,
            "modality": ("EA" if "EA" in (r.get("Intervention") or "") and
                         "TEAS" not in (r.get("Intervention") or "") else "TEAS"),
            "comparator_text": r.get("Comparator"),
            "comparator_category": comparator_category(r.get("Comparator")),
            "n_intervention": r.get("n I"), "n_comparator": r.get("n C"),
            "outcome_reported": r.get("Outcome"),
            "time_window": r.get("Timepoint"),
            "original_unit": r.get("Unit"),
            "reporting_format": r.get("Data type"),
            "value_intervention": r.get("Mean/Median I"),
            "variance_intervention": r.get("SD/IQR I"),
            "value_comparator": r.get("Mean/Median C"),
            "variance_comparator": r.get("SD/IQR C"),
            "exact_0_24h": r.get("Exact 0-24 h?"),
            "transformation_required": c["transformation"],
            "assumptions_required": c["assumptions"],
            "mme_conversion_factor": c["mme_factor"],
            "mme_factor_source": c["mme_factor_source"],
            "tier_mme": c["tier_mme"],
            "smd_eligible": c["smd_eligible"],
            "eligible_main_S0": "Yes" if c["tier_mme"] == "A" else "No",
            "reason": c["reason"],
            "v32_primary_eligibility": r.get("Primary eligibility"),
            "v32_analysis_readiness": r.get("Analysis readiness"),
            "v32_pipeline_inc_primary": v32_prim,
            "v32_pipeline_inc_sens": v32_sens,
            "source_qc_issue": r.get("Source-QC issue"),
            "verified_against_pdf": c["verified"],
            "source_url": r.get("Source URL"),
        })

    OUT.mkdir(parents=True, exist_ok=True)
    cols = list(audit[0].keys())
    csv_path = OUT / "PRIMARY_OUTCOME_DERIVABILITY_AUDIT_v33.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for a in audit:
            w.writerow({k: ("" if v is None else v) for k, v in a.items()})

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Derivability_Audit_v33"
    ws_out.append(cols)
    for a in audit:
        ws_out.append(["" if a[c] is None else a[c] for c in cols])
    wb_out.save(OUT / "PRIMARY_OUTCOME_DERIVABILITY_AUDIT_v33.xlsx")

    print(f"rows: {len(audit)}")
    from collections import Counter
    print("\ntier_mme distribution (rows):")
    for k, v in sorted(Counter(a["tier_mme"] for a in audit).items()):
        print(f"  Tier {k}: {v}")
    print("\nby study (best tier per study):")
    order = {"A": 0, "B+C": 1, "B": 1, "C": 2, "D": 3, "E": 4, "F": 5}
    best = {}
    for a in audit:
        s = a["study"]
        if s not in best or order.get(a["tier_mme"], 9) < order.get(best[s]["tier_mme"], 9):
            best[s] = a
    for k, v in sorted(Counter(b["tier_mme"] for b in best.values()).items()):
        print(f"  Tier {k}: {v} studies")
    print(f"\nwrote {csv_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
