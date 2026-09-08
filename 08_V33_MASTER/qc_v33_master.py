#!/usr/bin/env python3
"""
Pre-analysis QC gate for the v33 master.

Every assertion the integration brief requires, run against the built workbook
rather than against my intentions. Exit 1 on any failure so the pipeline stops
before Stata is invoked.
"""

from __future__ import annotations

import hashlib
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
VER = ROOT / "TEAS EA Verification"
V32 = VER / "TEAS_EA_RECONCILED_MASTER_DATA_v32_FINAL_LOCK_READY.xlsx"
V33 = VER / "TEAS_EA_RECONCILED_MASTER_DATA_v33_FINAL_LOCK_READY.xlsx"
V32_SHA256 = "74fda7d176fae15af4aa5bbff318514f998adcf957a2f9cad67e8de55ab7b12f"

GREEN, RED, BOLD, RESET = "\033[92m", "\033[91m", "\033[1m", "\033[0m"
fails: list[str] = []
passes = 0


def check(name: str, ok: bool, detail: str = "") -> None:
    global passes
    if ok:
        passes += 1
        print(f"{GREEN}PASS{RESET} | {name}")
    else:
        fails.append(name)
        print(f"{RED}FAIL{RESET} | {name}")
        for line in (detail or "(no detail)").splitlines():
            print(f"       {RED}{line}{RESET}")


def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for c in iter(lambda: f.read(1 << 20), b""):
            h.update(c)
    return h.hexdigest()


def sheet_rows(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) if h is not None else "" for h in rows[0]]
    return [dict(zip(hdr, r)) for r in rows[1:]]


def main() -> int:
    print(f"{BOLD}v33 MASTER QC{RESET}\n")

    # ---- v32 immutability -------------------------------------------------
    check("v32 master is unchanged (sha256 matches the committed blob)",
          sha256(V32) == V32_SHA256,
          f"expected {V32_SHA256}\nfound    {sha256(V32)}")

    if not V33.exists():
        check("v33 master exists", False, f"{V33} not found")
        return 1

    wb33 = openpyxl.load_workbook(V33, data_only=True)
    wb32 = openpyxl.load_workbook(V32, data_only=True)

    od33 = sheet_rows(wb33["Outcome_Data"])
    od32 = sheet_rows(wb32["Outcome_Data"])
    sm33 = sheet_rows(wb33["Study_Master"])

    # ---- architecture preserved ------------------------------------------
    missing = [s for s in wb32.sheetnames if s not in wb33.sheetnames]
    check("All v32 sheets are preserved in v33", not missing,
          f"missing: {missing}")
    check("v33 adds the two required audit sheets",
          {"v33_Change_Log", "v33_Supplement_Reconciliation"} <= set(wb33.sheetnames),
          f"sheets: {wb33.sheetnames}")

    # ---- study count ------------------------------------------------------
    studies = [r["Canonical study"] for r in sm33 if r.get("Canonical study")]
    check("Canonical study count == 70", len(studies) == 70,
          f"found {len(studies)}")
    dupes = [s for s, n in Counter(studies).items() if n > 1]
    check("No duplicated canonical studies", not dupes, f"duplicates: {dupes}")

    # ---- v32 rows preserved verbatim --------------------------------------
    # Compared numerically: re-serialising a cached float through openpyxl can
    # shift the last bit or two of its decimal rendering (-3.3000000000000007
    # -> -3.300000000000001). That is representation, not a data change, so the
    # check asserts numeric identity to 1e-9 and reports the worst deviation
    # rather than demanding byte-equal strings.
    mismatches, worst = [], 0.0
    for i, (x, y) in enumerate(zip(od32, od33[:len(od32)])):
        for k in x:
            a_, b_ = x.get(k), y.get(k)
            if a_ == b_:
                continue
            try:
                d = abs(float(a_) - float(b_))
                worst = max(worst, d)
                if d > 1e-9:
                    mismatches.append(f"row {i} col {k!r}: {a_!r} vs {b_!r}")
            except (TypeError, ValueError):
                mismatches.append(f"row {i} col {k!r}: {a_!r} vs {b_!r}")
    check(f"v33 Outcome_Data is a strict superset of v32 (max numeric drift {worst:.2e})",
          not mismatches, "\n".join(mismatches[:10]))

    added = od33[len(od32):]
    check("v33 appended exactly the staged rows", len(added) == 18,
          f"appended {len(added)}")

    # ---- primary-outcome protection ---------------------------------------
    bad = [r["Comparison ID"] for r in added
           if str(r.get("Primary opioid meta-analysis eligible?")).strip().lower() != "no"]
    check("No appended row is flagged primary-opioid eligible", not bad,
          f"offending: {bad}")

    bad = [r["Comparison ID"] for r in added
           if str(r.get("Exact 0-24 h postoperative opioid?")).strip().lower() != "no"]
    check("No appended row claims to be exact 0-24 h cumulative opioid", not bad,
          f"offending: {bad}")

    # strict primary k must still be 7 in the locked opioid sheet
    op = sheet_rows(wb33["Stata_Opioid24_Primary"])
    incl = [r for r in op
            if str(r.get("provisional_primary_include", "")).strip() in ("1", "1.0", "True")]
    check("Stata_Opioid24_Primary still yields strict primary k = 7",
          len(incl) == 7, f"found {len(incl)}: {[r.get('study_unit') for r in incl]}")

    # ---- duplicate contrast detection -------------------------------------
    key = lambda r: (str(r.get("Canonical study")), str(r.get("Outcome family")),
                     str(r.get("Outcome/result")), str(r.get("Timepoint/window")),
                     str(r.get("Intervention arm")), str(r.get("Comparator arm")))
    keys = [key(r) for r in od33]
    dup = [k for k, n in Counter(keys).items() if n > 1]
    check("No duplicated outcome contrasts on the composite key", not dup,
          "\n".join(" | ".join(k) for k in dup[:6]))

    cids = [str(r.get("Comparison ID")) for r in added]
    check("Appended Comparison IDs are unique",
          len(cids) == len(set(cids)), f"dupes: {[c for c,n in Counter(cids).items() if n>1]}")
    clash = sorted(set(cids) & {str(r.get("Comparison ID")) for r in od32})
    check("Appended Comparison IDs do not collide with v32", not clash, f"clash: {clash}")

    # ---- statistic-type integrity ----------------------------------------
    for r in added:
        pass
    bad = []
    for r in added:
        dt = str(r.get("Data type") or "")
        has_mean = r.get("Mean intervention") not in (None, "")
        has_med = r.get("Median intervention") not in (None, "")
        has_ev = r.get("Events intervention") not in (None, "")
        if "Median" in dt and (has_mean or has_ev):
            bad.append(f"{r['Comparison ID']}: median row carries mean/event fields")
        if dt == "Events/total" and (has_mean or has_med):
            bad.append(f"{r['Comparison ID']}: binary row carries mean/median fields")
        if dt == "Mean/SD" and (has_med or has_ev):
            bad.append(f"{r['Comparison ID']}: mean row carries median/event fields")
    check("No silent median->mean or binary->continuous conversion", not bad,
          "\n".join(bad))

    # ---- binary denominators ---------------------------------------------
    bad = []
    for r in added:
        if str(r.get("Data type")) != "Events/total":
            continue
        for ev, n, side in ((r.get("Events intervention"), r.get("Analyzed n intervention"), "i"),
                            (r.get("Events comparator"), r.get("Analyzed n comparator"), "c")):
            if ev is None or n is None:
                bad.append(f"{r['Comparison ID']}: missing events/denominator ({side})")
            elif float(ev) > float(n):
                bad.append(f"{r['Comparison ID']}: events {ev} > denominator {n} ({side})")
    check("Binary rows have valid denominators (events <= n)", not bad, "\n".join(bad))

    # ---- arm Ns present ---------------------------------------------------
    bad = [r["Comparison ID"] for r in added
           if not r.get("Analyzed n intervention") or not r.get("Analyzed n comparator")]
    check("Every appended row carries analysed arm Ns", not bad, f"missing: {bad}")

    # ---- multi-arm shared controls ---------------------------------------
    zhu = [r for r in added if r["Canonical study"] == "Zhu 2022"]
    bad = [r["Comparison ID"] for r in zhu
           if not str(r.get("Shared-control / multi-arm issue", "")).startswith("Yes")]
    check("Zhu 2022 multi-arm rows are flagged as sharing a control", not bad,
          f"unflagged: {bad}")
    ctrl_ns = {r.get("Analyzed n comparator") for r in zhu}
    check("Zhu 2022 shared control uses one consistent denominator",
          len(ctrl_ns) == 1, f"denominators seen: {ctrl_ns}")

    # ---- no manufactured MME ---------------------------------------------
    bad = []
    for r in added:
        unit = str(r.get("Unit/scale") or "")
        fam = str(r.get("Outcome family") or "")
        if fam in ("Rescue opioid use", "Rescue analgesia") and re.search(r"\bMME\b", unit):
            bad.append(f"{r['Comparison ID']}: rescue outcome expressed in MME")
        if "administrations" in unit and re.search(r"mg|ug|µg", unit):
            bad.append(f"{r['Comparison ID']}: count unit mixed with a dose unit")
    check("No rescue count or incidence was converted into an opioid dose", not bad,
          "\n".join(bad))

    # ---- HOLD / CONFLICT rows are not analysis eligible --------------------
    bad = []
    for r in added:
        oth = str(r.get("Other analysis eligibility") or "")
        if oth.startswith(("HOLD", "CONFLICTED")) and "Eligible" in oth:
            bad.append(r["Comparison ID"])
    check("HOLD and CONFLICTED rows are not marked eligible", not bad, f"offending: {bad}")

    holds = [r["Comparison ID"] for r in added
             if str(r.get("Other analysis eligibility") or "").startswith(("HOLD", "CONFLICTED"))]
    check("The expected HOLD/CONFLICT rows are present", len(holds) == 3,
          f"found {len(holds)}: {holds}")

    # ---- provenance -------------------------------------------------------
    bad = [r["Comparison ID"] for r in added
           if not str(r.get("Source location") or "").strip()
           or not str(r.get("Source PDF URL") or "").strip()]
    check("Every appended row records a source PDF and location", not bad, f"missing: {bad}")

    # ---- no NaN / undefined ----------------------------------------------
    # An empty cell is legitimate -- a mean/SD row has no median fields. What
    # must never appear is a value that has been *rendered* as a broken number.
    bad = []
    for r in added:
        for k, v in r.items():
            if v is None or v == "":
                continue
            s = str(v).strip().lower()
            if s in ("nan", "undefined", "#n/a", "inf", "-inf", "null", "#value!", "#ref!"):
                bad.append(f"{r['Comparison ID']}.{k} = {v}")
    check("No NaN / undefined values in appended rows", not bad, "\n".join(bad))

    # The v32 formula columns must survive as materialised values, not None.
    empties = sum(1 for r in od33[:len(od32)]
                  if r.get("Derived effect measure") in (None, ""))
    orig_empties = sum(1 for r in od32 if r.get("Derived effect measure") in (None, ""))
    check("v32 formula-derived columns survived as values, not blanks",
          empties == orig_empties,
          f"v33 has {empties} blank 'Derived effect measure' cells vs {orig_empties} in v32; "
          "openpyxl discards cached formula results unless they are materialised")

    # ---- reconciliation sheet covers all 40 supplement rows ---------------
    sr = sheet_rows(wb33["v33_Supplement_Reconciliation"])
    check("Reconciliation sheet covers all 40 supplement rows", len(sr) == 40,
          f"found {len(sr)}")
    undisposed = [r["Supplement row"] for r in sr if not str(r.get("Disposition") or "").strip()]
    check("Every supplement row has a disposition", not undisposed, f"rows: {undisposed}")

    print("\n" + "=" * 74)
    if fails:
        print(f"{RED}{BOLD}FAILED{RESET}  {passes}/{passes+len(fails)} checks passed")
        for f in fails:
            print(f"  {RED}x{RESET} {f}")
        return 1
    print(f"{GREEN}{BOLD}ALL CHECKS PASSED{RESET}  {passes}/{passes}")
    print("v33 master is safe to build analysis datasets from.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
