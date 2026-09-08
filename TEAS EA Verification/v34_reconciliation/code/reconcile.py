#!/usr/bin/env python3
"""Reconcile values to JSON/CSV only. XLSX authoring is in build_workbook.mjs.

The supplied PDF audit is a secondary extraction source; this run does not claim
to have independently re-extracted the PDFs. No graph digitization or imputation.
"""
import csv, json, re, math, hashlib, shutil
from pathlib import Path
from collections import Counter, defaultdict
from difflib import SequenceMatcher
import openpyxl

BASE=Path(__file__).resolve().parents[1]
ROOT=BASE.parent.parent
V33=BASE.parent/'TEAS_EA_RECONCILED_MASTER_DATA_v33_FINAL_LOCK_READY.xlsx'
AUDIT=Path('/Users/ryan/Downloads/TEAS_EA_v33_CONSOLIDATED_SOURCE_PDF_GAPFILL_AUDIT.xlsx')
def extract(path):
    w=openpyxl.load_workbook(path,data_only=False)
    return {s.title:[list(r) for r in s.iter_rows(values_only=True)] for s in w}, {
        'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
        'sheets':{s.title:{'rows':s.max_row,'columns':s.max_column,'formulas':sum(c.data_type=='f' for row in s for c in row),'merges':[str(x) for x in s.merged_cells.ranges], 'tables':[{'name':t.name,'ref':t.ref} for t in s.tables.values()]} for s in w},
        'named_ranges':list(w.defined_names)}
v,vi=extract(V33); a,ai=extract(AUDIT)
assert vi['sha256']=='64ef683c58a1faa2bf408f415d03d96dd6622abf8ded24a72f43db9c47a82a2a'
json.dump({'v33':vi,'audit':ai},open(BASE/'inputs/inventory.json','w'),ensure_ascii=False,indent=2)
shutil.copy2(AUDIT,BASE/'inputs'/AUDIT.name)
def records(data): return [dict(zip(data[0],r)) for r in data[1:] if any(x is not None for x in r)]
def norm(x):
    s=str(x or '').lower().replace('–','-').replace('—','-').replace('µ','u').replace('μ','u')
    s=re.sub(r'(\d)\s*h(?:ours?)?\b',r'\1h',s)
    return re.sub(r'\s+',' ',s).strip()
def number(x):
    if isinstance(x,(int,float)): return x
    s=str(x or '').strip()
    return float(s) if re.fullmatch(r'[+-]?\d+(?:\.\d+)?',s) else None
def pair(x):
    s=str(x or '').replace('–','-').replace('—','-')
    m=re.fullmatch(r'\s*(\d+(?:\.\d+)?)\s*[,\-]\s*(\d+(?:\.\d+)?)\s*',s)
    return [float(t) for t in m.groups()] if m else [None,None]
def j(x):return json.dumps(x,ensure_ascii=False,separators=(',',':'))
def writecsv(path,rows,headers=None):
    headers=headers or (list(rows[0]) if rows else ['record_id'])
    with open(path,'w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=headers,extrasaction='ignore'); w.writeheader();w.writerows(rows)
def matrix(rows,headers=None):
    headers=headers or list(rows[0]);return [headers]+[[r.get(h) for h in headers] for r in rows]

OH=v['Outcome_Data'][0]
EXTRA=['V34 record ID','V34 audit row','V34 disposition','Raw statistic intervention','Raw dispersion intervention','Raw statistic comparator','Raw dispersion comparator','Minimum intervention','Maximum intervention','Minimum comparator','Maximum comparator','IQR width intervention','IQR width comparator','Analysis population','V34 source filename','V34 source quote','V34 eligibility','V34 eligibility reason','V34 RoB2 status','V34 endpoint class','V34 modality','V34 comparator class','V34 time class']
old=records(v['Outcome_Data']); outcomes=[dict(r) for r in old]
src={r['study']:r for r in records(a['Source_Verification'])}
aud=records(a['All_Gapfill']); assert len(aud)==374
CANON={r[0] for r in v['Study_Master'][1:]};assert len(CANON)==70
for i,r in enumerate(outcomes,2):r['V34 record ID']=f'V33-OD-{i:04d}'
correction_reason={}; correction_type={}
def correct(excelrow,reason,kind='UPDATED_EXISTING',**values):
    r=outcomes[excelrow-2];r.update(values)
    correction_reason[r['V34 record ID']]=reason;correction_type[r['V34 record ID']]=kind
def note(excelrow,txt):
    r=outcomes[excelrow-2]; correct(excelrow,txt,**{'Source-QC issue':str(r['Source-QC issue'] or '')+' | V34: '+txt})

for er in range(179,184):
    assert outcomes[er-2]['Canonical study']=='Liang 2021'
    correct(er,'Source CONSORT: TEAS randomized 37, control 38; analyzed 35/35 unchanged.','METADATA_CORRECTION',**{'Randomized n intervention':37,'Randomized n comparator':38})
note(182,'SOURCE CONFLICT: abstract 48-h CRBD 4/35 versus 7/35; Table 2 4/35 versus 19/35. Retain table values; quantitative HOLD pending adjudication.')
correct(383,'Undefined analgesia metric retained in raw fields, removed from mean/SD fields; not opioid dose.',**{
    'Mean intervention':None,'SD intervention':None,'Mean comparator':None,'SD comparator':None,
    'Raw statistic intervention':'3.8','Raw dispersion intervention':'1.9','Raw statistic comparator':'5.0','Raw dispersion comparator':'2.9',
    'Derived effect measure':None,'Derived effect':None,'Other analysis eligibility':'AMBIGUOUS_HOLD: drug, unit, statistic, denominator and window undefined',
    'Record status':'AMBIGUOUS_HOLD / ORIGINAL VALUES PRESERVED IN RAW FIELDS'})
for er in range(225,229):
    vals={'Outcome family':'PCIA solution volume','Unit/scale':'mL multimodal PCIA solution','Exact 0-24 h postoperative opioid?':'No','Primary opioid meta-analysis eligible?':'No',
          'Other analysis eligibility':'Separate multimodal PCIA volume endpoint; never primary opioid mass'}
    if er==228: vals.update({'Outcome/result':'Cumulative multimodal PCIA solution volume','Data type':'Graph-only continuous data','Mean intervention':None,'SD intervention':None,'Mean comparator':None,'SD comparator':None,'Derived effect measure':None,'Derived effect':None})
    correct(er,'Figure 4 reports multimodal PCIA solution volume; no exact 24-h opioid mass is printed.',**vals)
    note(er,'Solution contains sufentanil, dexmedetomidine, flurbiprofen and palonosetron. No v34 conversion to opioid mass.')
assert outcomes[97-2]['Comparison ID']=='TU24_TEAS_vs_SHAM_TRAMADOL'
correct(97,'Audit Reconciliation_Status, Tu Table 4: tramadol use is 6–24 h, not cumulative 0–24 h.','METADATA_CORRECTION',**{'Timepoint/window':'6–24 h','Source location':'Table 4 / rescue tramadol, 6–24 h','Other analysis eligibility':'Separate 6–24 h binary rescue-opioid result; excluded from cumulative 0–24 h rescue set'})
note(94,'Table 3/abstract P=0.047; Results prose P=0.47; retain table P, preserve both.')
for er in (367,368):note(er,'Yu Table 2 TEAS SDs POD1/POD2 = 1.53/0.98; abstract = 1.41/0.88. POD2 prose P=0.26, table *P<0.05, abstract significant. Table values retained; pain quantitative HOLD.')
note(147,'Long Table 7 EA 24-h VAS 2.78±0.97; abstract 2.65±0.94. Table values retained; quantitative HOLD pending adjudication.')
for er in (145,146):note(er,'PND denominator conflict: study-level n=27/26 differs from percentages implying 28/25. No denominator reconstruction; HOLD.')
correct(109,'Sim Results reports median with IQR width, not Q1/Q3.','METADATA_CORRECTION',**{'Data type':'Median/IQR width as printed','IQR width intervention':0.05,'IQR width comparator':0.11,'Raw dispersion intervention':'IQR 0.05','Raw dispersion comparator':'IQR 0.11','Q1 intervention':None,'Q3 intervention':None,'Q1 comparator':None,'Q3 comparator':None})
note(105,'Historical aggregate placeholder superseded for active analysis by three audit arm-specific converted-metric rows; no double counting. Author-defined intraoperative metric is not review postoperative MME.')
correct(105,'Aggregate placeholder retained for history but superseded by arm-specific audit records.',**{'Record status':'HISTORICAL_AGGREGATE_SUPERSEDED','Other analysis eligibility':'NO: superseded aggregate placeholder; retain for provenance'})

# Trial-level allocation correction also applied to all Zhu original contrasts.
for er,r in enumerate(outcomes,2):
    if r['Canonical study']=='Zhu 2022':
        correct(er,'Zhu CONSORT allocation 103/103/103/104 retained over conflicting abstract.','METADATA_CORRECTION',**{'Randomized n intervention':103,'Randomized n comparator':104})

# The full study-level comparison was reviewed before defining this alias.
# Audit Li 6-h VAS>=4 exactly matches v33 30/140 vs 50/140; later clock points differ.
EXACT_ALIASES={138:155}
HOLD_AUDIT=set([20,21,22,74,95,96,98,102,104,114,130,156,174,175,176,179,298,299,310,357,365,366,367,368,369,370,371,372])
# Per-protocol subsets conditioned on developing chronic pain need estimand adjudication.
HOLD_AUDIT.update(range(302,306))
HOLD_OLD={147,145,146,182,367,368,383}

dispositions=[]; comparisons=[]
for er,x in enumerate(aud,2):
    st=x['study'];assert st in CANON
    r={h:None for h in OH+EXTRA}
    maps={'Canonical study':'study','Intervention arm':'intervention_arm','Comparator arm':'comparator_arm','Outcome family':'outcome_family','Outcome/result':'outcome_result','Timepoint/window':'timepoint_window','Data type':'statistic_type','Unit/scale':'unit_scale','Reported P':'reported_p','Analysis population':'analysis_population','V34 source quote':'verbatim_quote','Source-QC issue':'ambiguity_notes'}
    for h,k in maps.items():r[h]=x[k]
    for h,k in [('Randomized n intervention','randomized_n_int'),('Randomized n comparator','randomized_n_comp'),('Analyzed n intervention','analyzed_n_int'),('Analyzed n comparator','analyzed_n_comp')]:r[h]=number(x[k])
    for side,suf in [('intervention','int'),('comparator','comp')]:
        r[f'Raw statistic {side}']=x[f'value_{suf}'];r[f'Raw dispersion {side}']=x[f'dispersion_{suf}']
        stat=x['statistic_type']
        if stat.startswith('Mean'):
            r[f'Mean {side}']=number(x[f'value_{suf}']);r[f'SD {side}']=number(x[f'dispersion_{suf}'])
        elif stat.startswith('Median'):
            r[f'Median {side}']=number(x[f'value_{suf}'])
            if stat=='Median/IQR':r[f'Q1 {side}'],r[f'Q3 {side}']=pair(x[f'dispersion_{suf}'])
            elif stat=='Median/range':r[f'Minimum {side}'],r[f'Maximum {side}']=pair(x[f'dispersion_{suf}'])
            elif 'width' in stat:r[f'IQR width {side}']=number(str(x[f'dispersion_{suf}']).replace('IQR','').strip())
        if stat.startswith('Events') or stat.startswith('Published event count'):
            r[f'Events {side}']=number(x[f'events_{suf}'])
    r.update({'V34 record ID':f'AUDIT-{er:04d}','V34 audit row':er,'Comparison ID':f'V34_B{x["batch"]}_{x["batch_row"]:03d}',
              'V34 source filename':src[st]['source_pdf_filename'],'Source PDF URL':src[st]['drive_link'],
              'Source location':f'{x["page"]}; {x["table_or_figure"]}','Directly reported?':'Yes','Derived/converted?':'No',
              'Exact 0-24 h postoperative opioid?':'No','Primary opioid meta-analysis eligible?':'No',
              'Shared-control / multi-arm issue':'Yes — shared comparator; select one contrast per model' if st in ['Zhu 2022','Sun 2017','Sim 2002','Lu 2021'] else 'No (preserve trial strata where specified)',
              'Record status':'SOURCE-AUDIT RECONCILED; supplied audit, no independent PDF reread'})
    if er in range(242,245):r['Derived/converted?']='Yes — author-defined intraoperative conversion; no v34 conversion'
    # Robust missing-data distinction uses the result fields, never a generic NR note.
    raw=[x[k] for k in ['value_int','value_comp','events_int','events_comp']]
    if 'graph' in norm(stat):disp='GRAPH_ONLY_ADDED'
    elif all(str(y)=='NOT REPORTED' for y in raw):disp='NOT_REPORTED_ADDED'
    elif er==310:disp='CONFLICT_HOLD';r['Analyzed n intervention']=None;r['Analyzed n comparator']=None
    elif er in HOLD_AUDIT:disp='AMBIGUOUS_HOLD'
    elif stat.startswith('Median') or stat.startswith('Published percentage') or er in range(242,245) or er==319:disp='NOT_ANALYSIS_ELIGIBLE'
    elif stat=='Other (describe)':disp='AMBIGUOUS_HOLD'
    else:disp='ADDED'
    if er in EXACT_ALIASES:
        tgt=outcomes[EXACT_ALIASES[er]-2]
        assert r['Events intervention']==tgt['Events intervention'] and r['Events comparator']==tgt['Events comparator']
        assert r['Analyzed n intervention']==tgt['Analyzed n intervention'] and r['Analyzed n comparator']==tgt['Analyzed n comparator']
        disp='ALREADY_PRESENT'; rid=tgt['V34 record ID']
    else:
        rid=r['V34 record ID'];r['V34 disposition']=disp;outcomes.append(r)
    # Retain transparent candidate comparisons for every row, including Batch 1.
    candidates=[(i,z) for i,z in enumerate(old,2) if z['Canonical study']==st]
    best=sorted(candidates,key=lambda iz:SequenceMatcher(None,norm(x['outcome_result'])+' '+norm(x['timepoint_window']),norm(iz[1]['Outcome/result'])+' '+norm(iz[1]['Timepoint/window'])).ratio(),reverse=True)[:3]
    comparisons.append({'audit_row':er,'study':st,'v33_study_rows_compared':len(candidates),'closest_v33_rows':';'.join(str(i) for i,z in best),'disposition':disp,'match_v33_row':EXACT_ALIASES.get(er),'rationale':'Exact semantic and numerical match; retain existing row' if er in EXACT_ALIASES else 'Study-level review: distinct result, time window, arm, distribution or source-unavailability record; raw source preserved. Closest text matches are candidates, not automatic merges.'})
    dispositions.append({'audit_row':er,'batch':x['batch'],'batch_row':x['batch_row'],'study':st,'disposition':disp,'v34_record_id':rid,'v33_excel_row':EXACT_ALIASES.get(er),'outcome':x['outcome_result'],'window':x['timepoint_window']})

# Explicit supplement-dependent outcomes: two records, not generic invented endpoints.
for n,title,window in [(1,'Total hospital length of stay','Total admission'),(2,'30-day postoperative complications','0–30 days')]:
    r={h:None for h in OH+EXTRA};r.update({'Canonical study':'Gao 2021','Comparison ID':f'V34_GAO21_SUPP_{n}','Intervention arm':'TEAS','Comparator arm':'Sham','Outcome family':'Length of stay' if n==1 else 'Adverse events','Outcome/result':title,'Timepoint/window':window,'Data type':'Source not accessed','V34 record ID':f'SUPP-GAO-{n}','V34 disposition':'SOURCE_NOT_ACCESSED','Source location':'Appendix 2 / Supplementary Table S4','Source PDF URL':src['Gao 2021']['drive_link'],'V34 source filename':src['Gao 2021']['source_pdf_filename'],'Source-QC issue':'Main paper directs exact values to unaccessed supplement. Not equivalent to NOT REPORTED.','Record status':'SOURCE NOT ACCESSED','Directly reported?':'No — supplement not accessed','Derived/converted?':'No','Exact 0-24 h postoperative opioid?':'No','Primary opioid meta-analysis eligible?':'No'})
    outcomes.append(r)

PRIMARY_IDS={r['comparison_id'] for r in records(v['Stata_Opioid24_Primary']) if r['provisional_primary_include']==1}
assert len(PRIMARY_IDS)==7
oldready={r['Comparison ID']:r for r in records(v['Analysis_Readiness'])}
rob={r['Canonical study']:r for r in records(v['Corrected_RoB2'])}
def endpoint(r):
    s=norm(r['Outcome/result']);f=norm(r['Outcome family']);u=norm(r['Unit/scale'])
    if r.get('V34 audit row') in range(115,120):return 'QoR-40 subscale'
    if 'intraoperative' in s+' '+f:return 'Intraoperative opioid' if any(t in s+' '+f+' '+u for t in ['opioid','fentanil','morphine']) else 'Intraoperative anesthetic'
    if 'solution' in u+s or 'pcia solution volume' in f:return 'PCIA solution volume'
    if 'time to first' in s and any(t in s for t in ['rescue','analgesi']):return 'Time to first rescue'
    if any(t in s for t in ['press','attempt','demand']) and any(t in s for t in ['pca','pcia','pump']):return 'PCA/PCIA button presses'
    if any(t in s for t in ['deliveries','delivery','boluses']) and any(t in s for t in ['pca','pcia']):return 'PCA deliveries / bolus count'
    if 'analgesia requirement' in s and r['Canonical study']=='Liang 2021':return 'Ambiguous analgesia requirement'
    if any(t in s for t in ['antiemetic','metoclopramide','tropisetron']):return 'Rescue antiemetic'
    if any(t in s for t in ['rescue','analgesic drug','analgesia use','analgesic use']):
        if any(t in s for t in ['flurbiprofen','parecoxib','ketorolac','dexketoprofen','nsaid']):return 'Non-opioid rescue analgesia'
        if any(t in s for t in ['sufentanil','tramadol','morphine','dezocine','bucinnazine','opioid','pethidine']):return 'Rescue opioid use' if r['Events intervention'] is not None else 'Rescue opioid frequency / dose (native unit)'
        return 'Rescue analgesia — drug class unresolved'
    if 'qor-15' in s+u:return 'QoR-15'
    if 'qor-40' in s+u:return 'QoR-40 subscale' if any(x in s for x in ['physical comfort','emotional','physical independence','psychological support','pain subscale']) else 'QoR-40'
    if 'vomit' in s and 'nausea' not in s:return 'Vomiting severity' if any(t in s for t in ['severity','grade','score']) else 'Vomiting'
    if 'nausea' in s and 'vomit' not in s:return 'Nausea severity' if any(t in s for t in ['severity','grade','score']) else 'Nausea'
    if 'ponv' in s or ('nausea' in s and 'vomit' in s):return 'PONV severity' if any(t in s for t in ['severity','grade','score']) else 'PONV composite'
    if f=='pain' or ('pain' in s and f!='chronic pain'):
        context='rest' if 'rest' in s else 'cough' if 'cough' in s else 'movement/activity' if any(t in s for t in ['movement','activity','mobile','dynamic']) else 'unspecified context'
        return 'Pain — '+context
    if 'flatus' in s:return 'GI — first flatus'
    if 'defecation' in s or 'bowel motion' in s:return 'GI — first defecation'
    if 'bowel sound' in s or 'bowel-sound' in s:return 'GI — bowel sounds'
    if 'ileus' in s or 'bowel obstruction' in s:return 'GI — ileus/obstruction (separate definitions)'
    if 'distention' in s or 'distension' in s:return 'GI — abdominal distention'
    if 'water' in s or 'oral intake' in s:return 'GI — oral intake'
    if any(t in s for t in ['solid food','normal diet','soft diet']):return 'GI — diet tolerance'
    if 'pacu' in s or 'recovery-room' in s:return 'Functional recovery — PACU'
    if any(t in s for t in ['length of stay','hospital stay','postoperative stay']) or f in ('length of stay','los'):return 'LOS'
    if 'opioid' in f:return 'Postoperative opioid dose (native window/unit)'
    return r['Outcome family'] or 'Unclassified'
def modality(r):
    if r['Canonical study'] in ['Zhu 2022','Yang 2020','Huang 2025','Long 2025','Sim 2002']:return 'EA'
    if r['Canonical study'] in src:return 'TEAS' # remaining audited interventions are explicitly transcutaneous
    s=norm(r['Intervention arm'])
    if any(t in s for t in ['teas','tens','transcutaneous']):return 'TEAS'
    if re.search(r'\bea\b',s) or 'electroacupuncture' in s:return 'EA'
    return 'MODALITY_REVIEW_REQUIRED'
def comparator(r):
    s=norm(r['Comparator arm'])
    if r['Canonical study']=='Pan 2023':return 'Usual care' # existing source-verified comparison is Control/usual care
    if r['Canonical study']=='Lu 2022':return 'Sham' # existing source: No-current control + ERAS
    if any(t in s for t in ['sham','placebo','no current','no-current','no stimulation','0 ma','zero-current']):return 'Sham'
    if any(t in s for t in ['usual','standard','alone','no acupuncture','no-acupuncture','no treatment','uc group']):return 'Usual care'
    return 'CONTROL_TYPE_REVIEW_REQUIRED'
def timeclass(r):
    t=norm(r['Timepoint/window'])
    if r['Comparison ID'] in PRIMARY_IDS:return 'cumulative 0–24 h (locked primary source adjudication)'
    if any(x in t for x in ['ambiguous','unclear','unspecified','not stated']):return 'AMBIGUOUS'
    if 'pod' in t or 'postoperative day' in t or re.search(r'day\s*\d',t):return t
    if any(x in t for x in ['mean','average']):return t
    if re.search(r'\d\s*-\s*\d',t):return t
    if re.fullmatch(r'(?:at )?24h(?: (?:after (?:surgery|operation)|postoperatively|postoperative|after the operation))?',t):return '24 h point'
    return t
def eligibility(r):
    er=r.get('V34 audit row');dt=norm(r['Data type']);status=norm(r['Record status']);f=r['V34 endpoint class'];tp=r['V34 time class']
    if r['Comparison ID'] in PRIMARY_IDS:return 'INCLUDE','Locked strict primary result; prespecified source adjudication and conversions retained.'
    if 'source not accessed' in dt+' '+status:return 'SOURCE NOT ACCESSED','Referenced supplement not accessed; no result values populated.'
    if 'graph' in dt:return 'GRAPH ONLY' if 'digitized' not in dt else 'CONDITIONAL','No new graph digitization; historical digitized data require their original validation before pooling.'
    if er in HOLD_AUDIT and er!=357:return 'HOLD','Source definition, timepoint, denominator, scale, statistic, or post-randomization subset requires adjudication; see raw source/QC.'
    if r['V34 record ID'] in {f'V33-OD-{x:04d}' for x in HOLD_OLD}:return 'HOLD','Unresolved source conflict or undefined endpoint; printed values retained, excluded from quantitative models.'
    if r.get('V34 disposition')=='NOT_REPORTED_ADDED' or 'not reported' in dt or 'not applicable' in dt:return 'NOT REPORTED','Accessed source does not supply the requested numerical result.'
    if 'superseded' in status:return 'EXCLUDE','Historical aggregate placeholder superseded by arm-specific rows; not an independent result.'
    if r['Canonical study'].startswith('Yeh '):return 'DUPLICATE HOLD','Historical possible overlapping publication family; retain evidence, exclude independent pooling.'
    if 'denominator inconsistent' in dt or 'ambiguous' in dt or tp=='AMBIGUOUS':return 'HOLD','Unresolved source statistic, denominator or timepoint.'
    if r['Canonical study']=='Wu 2025' and f=='Intraoperative opioid':return 'EXCLUDE','Intraoperative dose precedes PACU randomization/intervention; baseline covariate.'
    if er==319:return 'EXCLUDE','Postoperative EA cannot estimate a treatment effect on preceding intraoperative alfentanil.'
    if 'median' in dt:return 'NARRATIVE ONLY','Median and original dispersion preserved; no mean/SD conversion or MD pooling.'
    if 'percentage' in dt:return 'NARRATIVE ONLY','Printed percentage/CI retained; event counts not reverse-engineered.'
    if er in range(242,245) or 'study-defined converted' in dt:return 'CONDITIONAL','Author-defined intraoperative morphine-equivalent metric; no review MME conversion, requires conversion-method adjudication.'
    if dt!='mean/sd' and dt!='events/total':return 'HOLD','Nonstandard or incomplete statistic; no inferred quantities.'
    ns=[number(r['Analyzed n intervention']),number(r['Analyzed n comparator'])]
    if None in ns or min(ns)<=0:return 'HOLD','Missing or invalid result-specific denominators.'
    if dt=='mean/sd':
        vals=[number(r[k]) for k in ['Mean intervention','SD intervention','Mean comparator','SD comparator']]
        if None in vals:return 'HOLD','Mean or SD not available; no imputation.'
        if min(vals[1],vals[3])<0:return 'HOLD','Invalid negative SD.'
    else:
        ev=[number(r['Events intervention']),number(r['Events comparator'])]
        if None in ev or any(e<0 or e>n or e!=int(e) for e,n in zip(ev,ns)):return 'HOLD','Invalid or missing event/denominator pair.'
    if any(t in norm(r['Unit/scale']) for t in ['undefined','not stated','not specified','ambiguous']):return 'HOLD','Scale/measurement definition unresolved.'
    prev=oldready.get(r['Comparison ID'],{})
    if prev.get('Analysis readiness') in ['QC HOLD','SOURCE-QC HOLD','EXCLUDE']:
        return 'HOLD','Preserved v33 source-QC gate: '+str(prev.get('Readiness reason'))
    return 'INCLUDE','Numerically extractable in its native endpoint/time/statistic; pooling still requires compatible stratum, comparator and independent-arm selection.'

for r in outcomes:
    r['V34 endpoint class']=endpoint(r);r['V34 modality']=modality(r);r['V34 comparator class']=comparator(r);r['V34 time class']=timeclass(r)
    e,reason=eligibility(r);r['V34 eligibility']=e;r['V34 eligibility reason']=reason
    r['V34 RoB2 status']='ROB2_RESULT_SPECIFIC_PENDING'
    # Only exact original result/time matches are linked; new outcomes never inherit study-wide RoB.
    prev=oldready.get(r['Comparison ID'],{})
    if r.get('V34 audit row') is None and str(prev.get('Additional outcome-specific RoB needed?','')).startswith('No'):
        r['V34 RoB2 status']='EXISTING_RESULT_SPECIFIC: '+str(prev.get('Selected RoB2 overall judgment'))
    if r['Comparison ID'] in PRIMARY_IDS:r['V34 RoB2 status']='EXISTING_LOCKED_PRIMARY: '+str(next(x['rob_overall'] for x in records(v['Stata_Opioid24_Primary']) if x['comparison_id']==r['Comparison ID']))
    if r.get('V34 audit row') or r.get('V34 disposition')=='SOURCE_NOT_ACCESSED':r['Other analysis eligibility']=e+' — '+reason

assert len({r['V34 record ID'] for r in outcomes})==len(outcomes)
assert {r['Canonical study'] for r in old}=={r['Canonical study'] for r in outcomes}
byid={r['Comparison ID']:r for r in outcomes}

# Machine-check exact composite keys with non-superficial distinctions retained.
KEY_FIELDS=['Canonical study','Outcome family','Outcome/result','Timepoint/window','Intervention arm','Comparator arm','Data type','Analyzed n intervention','Analyzed n comparator','Unit/scale']
def key(r):return tuple(norm(r.get(k)) for k in KEY_FIELDS)
keys=defaultdict(list)
for r in outcomes:keys[key(r)].append(r['V34 record ID'])
dups=[ids for ids in keys.values() if len(ids)>1]
oldkeys=Counter(key(r) for r in old)
introduced=[ids for k,ids in keys.items() if len(ids)>max(1,oldkeys[k])]
assert not introduced, introduced

# Conflict table includes explicit competing source values, not just hidden notes.
CH=['study','outcome','timepoint','source_location_A','value_A','source_location_B','value_B','current_structured_choice','reason_for_choice','analysis_status','requires_author_contact','notes']
conf=[]
def conflict(st,o,t,loca,va,locb,vb,choice,reason,status='HOLD',contact='Yes',notes=''):
    conf.append(dict(zip(CH,[st,o,t,loca,va,locb,vb,choice,reason,status,contact,notes])))
conflict('Liang 2021','CRBD','48 h','Abstract','4/35 vs 7/35','Table 2','4/35 vs 19/35','4/35 vs 19/35','Retain v33 table hierarchy; do not erase abstract conflict')
conflict('Liang 2021','Randomized arm counts','Allocation','v33 CRBD rows','38/37','Source Figure 2','37/38','37/38','Correct to CONSORT; analyzed 35/35 unchanged','METADATA_CORRECTED','No')
conflict('Liang 2021','Extra analgesia requirement','Unclear','Table 4','3.8 (1.9) vs 5.0 (2.9)','Endpoint definition','Drug, unit, measurement, denominator and window undefined','Raw values only','Do not treat as dose/count')
conflict('Liang 2021','PONV definition','End of surgery/PACU/24h/48h','Table 4','Row labelled PONV; printed event counts retained','Methods','Nausea verbal scale described; composite event definition unclear','Source label and counts retained; quantitative HOLD','Clarify whether these represent nausea or composite PONV; no silent relabelling')
for day,ta,ab in [(1,'3.70±1.53 vs 4.73±1.53','TEAS 3.70±1.41'),(2,'1.83±0.98 vs 2.30±0.95','TEAS 1.83±0.88')]:conflict('Yu 2020','Resting VAS dispersion',f'POD{day}','Table 2',ta,'Abstract',ab,ta,'Table hierarchy retained; no SD substitution')
conflict('Yu 2020','Pain significance','POD2','Results prose','P=0.26','Table 2 / abstract','*P<0.05 / significant','Both P versions','P is not repaired by assumption')
conflict('Yu 2020','MMSE direction','POD1/POD2','Table 3','TEAS numerically higher','Results prose','TEAS lower','Table values','Direction conflict preserved','SOURCE_NOTE / NUMERIC_CANDIDATE')
for o,pt,pp in [('Rest NRS','0.94','0.12'),('Cough NRS','0.89','0.49'),('Rescue parecoxib','0.35','0.58')]:conflict('Lu 2021',o,'24 h','Table 3','P='+pt,'Results prose','P='+pp,'Table values; both P versions','P conflict does not change extracted group values','P_CONFLICT_RETAINED')
conflict('Lu 2022','QoR-15','24 h','Table 3','Median 54.5 (range 15–92) vs 59 (27–136), P=0.01','Results prose','Median 50.5 (IQR 42.3–61.8) vs 44.5 (35–57), P=0.04','Table 3','Preserve conflicting statistic, magnitude and direction; no averaging')
conflict('Long 2025','VAS','24 h','Table 7','EA 2.78±0.97','Abstract','EA 2.65±0.94','Table 7','Table hierarchy retained, material numeric conflict')
conflict('Long 2025','PND denominator','POD1/POD3/POD7','Study flow','EA 27 / C 26','Table 2 count/percentage pairs','Imply EA 28 / C 25; POD7 1 (3.6%) vs 2 (8.0%)','Printed events; no assumed result denominator','No denominator reconstruction')
conflict('Zheng 2025','Pain significance','24 h','Results prose','P=0.042','Table 4 footnote','P=0.422','Both versions','Numerical means/SD unchanged; no invented P correction','P_CONFLICT_RETAINED')
conflict('Gu 2019','PCIA quantity','24 h','v33 label','Exact cumulative sufentanil','Figure 4','Graph-only multimodal PCIA solution volume','Graph-only mL solution','Reclassify; no pure opioid mass','CORRECTED / GRAPH_ONLY','No')
conflict('Gu 2019','Satisfaction','Postoperative','Prose','Composite 67.3% vs 42.4%','Table 4','Category counts: 32/9, 19/32, 7/16, 0/2','Four printed category contrasts','Do not reconstruct prose composite','CATEGORY_DATA_ONLY')
conflict('Zhu 2022','Randomized allocation','Baseline','Abstract','Pre103 / 30min104 / Comb103 / Usual103','CONSORT Figure 1','Pre103 / 30min103 / Comb103 / Usual104','CONSORT','Prespecified source flow hierarchy','METADATA_CORRECTED','Yes')
conflict('Zhu 2022','Movement NRS','6–24 h','Methods','Movement pain at 24 h','Table 3','Pain at 6–24 h','6–24 h interval','Do not promote interval to exact clock point','INTERVAL_ONLY')
conflict('Zhu 2022','QoR-15','AMBIGUOUS','Table 3','Printed QoR-15 means/SD','Assessment label','Time not explicitly identified','AMBIGUOUS','No assumed 24-h window')
conflict('Jiang 2026','NRS','AMBIGUOUS','Table 4','3 (2–3) vs 3 (2–3.5)','Methods','6,24,48 h planned; result not linked','AMBIGUOUS','No assumed 24-h window')
conflict('Jiang 2026','Printed effect direction/CI','Defecation/cost','Table 4 / Results','Raw estimates and CI retained in V34_Audit_Dispositions','Arm summaries','Effect direction/CI formatting inconsistent','Group summaries unchanged; no corrected effect asserted','Use native data only if compatible; preserve printed effect','SOURCE_NOTE')
conflict('Tu 2024','Vomiting P','2–6 h','Table 3 / abstract','0.047','Results prose','0.47','0.047','Table hierarchy; preserve both','P_CONFLICT_RETAINED')
conflict('Tu 2024','Rescue tramadol window','6–24 h','v33 Outcome_Data','Within 24 h','Audit Reconciliation_Status / Table 4','6–24 h','6–24 h','Exclude from cumulative 0–24 h rescue set','CORRECTED','No')
conflict('Pan 2023','Participant flow','Trial flow','Methods','120 women','Abstract / Results','105 final cohort, 52/53','Analyzed 52/53; no inferred randomized allocation','Source flow inconsistency','METADATA_HOLD')
conflict('Pan 2023','Mobile/PACU NRS statistic','1h/24h/PACU','Tables 6–7','Parenthetical numbers, e.g. 2 (1.4) vs 3 (2.4)','Statistic definition','Parenthetical statistic undefined','Raw values only','No SD/IQR assumption')
conflict('Huang 2025','QoR-40 direction','Discharge','Table 4','Median 196 vs 192','Results prose','Lower with EA','Table medians','No conversion; direction conflict retained','NARRATIVE_ONLY')

updates={}
updates['Outcome_Data']=matrix(outcomes,OH+EXTRA)
RH=v['Analysis_Readiness'][0]+EXTRA
ready=[]
for r in outcomes:
    rr=dict(r);rr.update({'Normalized synthesis family':r['V34 endpoint class'],'Timepoint class':r['V34 time class'],'Analysis readiness':r['V34 eligibility'],'Readiness reason':r['V34 eligibility reason'],'Preferred candidate within family?':'Model-specific; see V34_Secondary_Ready','Duplicate/overlap cluster':r['Canonical study'] if 'Yes' in str(r['Shared-control / multi-arm issue']) or r['V34 eligibility']=='DUPLICATE HOLD' else None,'Conversion/handling needed':'No new imputation, digitization, median conversion or proxy MME','Selected RoB2 overall judgment':r['V34 RoB2 status'],'Additional outcome-specific RoB needed?':'Yes' if 'PENDING' in r['V34 RoB2 status'] else 'No — exact existing result link','Independent randomized-study unit':r['Canonical study'],'Recommended comparison role':'Locked strict primary' if r['Comparison ID'] in PRIMARY_IDS else 'Native secondary candidate / hold; see status','Independent-study counting note':'One independent contrast per study per compatible model; preserve multi-arm/stratum dependencies.'})
    ready.append(rr)
updates['Analysis_Readiness']=matrix(ready,RH)
setrules={'Set_Opioid_24h':lambda r:r['Comparison ID'] in PRIMARY_IDS,
 'Set_Pain_24h':lambda r:r['V34 endpoint class'].startswith('Pain') and any(x in norm(r['Timepoint/window']) for x in ['24','pod1','day 1']),
 'Set_PONV':lambda r:any(r['V34 endpoint class'].startswith(x) for x in ['PONV','Nausea','Vomiting','Rescue antiemetic']),
 'Set_QoR':lambda r:r['V34 endpoint class'].startswith('QoR'),
 'Set_GI_Recovery':lambda r:r['V34 endpoint class'].startswith('GI')}
for name,rule in setrules.items():updates[name]=matrix([r for r in ready if rule(r)],v[name][0]+['V34 record ID','V34 endpoint class','V34 modality','V34 comparator class','V34 time class'])
# Explicit opioid candidates retain all native categories/holds, strict flag only for seven.
candidates=[]
for r in ready:
    if any(s in r['V34 endpoint class'].lower() for s in ['opioid','pca','pcia','analgesia']):
        candidates.append(dict(zip(v['Opioid_24h_Candidates'][0],[r['Canonical study'],r['Comparison ID'],r['Intervention arm'],r['Comparator arm'],r['Outcome/result'],r['Timepoint/window'],r['Data type'],r['Analyzed n intervention'],r['Analyzed n comparator'],r['Mean intervention'] if r['Mean intervention'] is not None else r['Median intervention'],r['SD intervention'] if r['SD intervention'] is not None else r['Raw dispersion intervention'] if 'Raw dispersion intervention' in r else None,r['Mean comparator'] if r['Mean comparator'] is not None else r['Median comparator'],r['SD comparator'],r['Unit/scale'],r['Exact 0-24 h postoperative opioid?'],'Yes' if r['Comparison ID'] in PRIMARY_IDS else 'No — separate candidate/hold',r['V34 eligibility'],r['V34 eligibility reason'],r['Duplicate/overlap cluster'],r['Source-QC issue'],r['Source PDF URL']])))
updates['Opioid_24h_Candidates']=matrix(candidates,v['Opioid_24h_Candidates'][0])

# Regenerate locked primary source values, never copy stale numeric exports.
native_map={'study_unit':'Canonical study','comparison_id':'Comparison ID','intervention':'Intervention arm','comparator':'Comparator arm','outcome':'Outcome/result','time_window':'Timepoint/window','data_type':'Data type','n_i':'Analyzed n intervention','n_c':'Analyzed n comparator','mean_i':'Mean intervention','sd_i':'SD intervention','median_i':'Median intervention','q1_i':'Q1 intervention','q3_i':'Q3 intervention','events_i':'Events intervention','mean_c':'Mean comparator','sd_c':'SD comparator','median_c':'Median comparator','q1_c':'Q1 comparator','q3_c':'Q3 comparator','events_c':'Events comparator','unit':'Unit/scale','source_qc':'Source-QC issue','source_url':'Source PDF URL'}
primary=[]
for z in records(v['Stata_Opioid24_Primary']):
    if z['comparison_id'] not in PRIMARY_IDS:continue
    r=byid[z['comparison_id']];p=dict(z)
    p.update({k:r[h] for k,h in native_map.items()});p['v34_record_id']=r['V34 record ID'];primary.append(p)
assert len(primary)==len({r['study_unit'] for r in primary})==7
updates['Stata_Opioid24_Primary']=matrix(primary)
writecsv(BASE/'data/v34_stata_opioid24_primary_native.csv',primary)
PLH=v['Opioid_Primary_StudyLevel'][0];pls=[]
for p in primary:
    r=byid[p['comparison_id']]
    pls.append(dict(zip(PLH,[r['Canonical study'],r['Comparison ID'],r['Intervention arm'],r['Comparator arm'],r['Outcome/result'],r['Timepoint/window'],r['Data type'],r['Analyzed n intervention'],r['Analyzed n comparator'],f'{r["Mean intervention"]} ± {r["SD intervention"]}',f'{r["Mean comparator"]} ± {r["SD comparator"]}',r['Unit/scale'],'Strict primary','INCLUDE IN PRIMARY','Preserved locked source adjudication',p['conversion_needed'],p['rob_overall'],p['sensitivity_flag'],p['duplicate_multiarm_handling'],r['Source-QC issue'],r['Source PDF URL']])))
updates['Opioid_Primary_StudyLevel']=matrix(pls,PLH)
updates['Opioid_Primary_Summary']=[['24-h continuous opioid study-level set','Count'],['Primary study units',len(primary)],['TEAS / sham',sum(modality(byid[p['comparison_id']])=='TEAS' for p in primary)],['EA / usual care',sum(modality(byid[p['comparison_id']])=='EA' for p in primary)],['Total analyzed participants',sum(p['n_i']+p['n_c'] for p in primary)],['Primary direct set excluding El-Rakshy high-RoB sensitivity case',sum(p['study_unit']!='El-Rakshy 2009' for p in primary)]]

# Preserve old AF sheets as history; active source layer propagates exact-ID corrections.
updates['V33_Outcome_Data_AF_LOCK']=v['Outcome_Data_AF_LOCK']
updates['V33_Stata_AF_Long']=v['Stata_AF_Long']
af=[]
afids=set()
for r in records(v['Outcome_Data_AF_LOCK']):
    z=dict(r);rid=r['Comparison ID'];afids.add(rid)
    # Existing AF source-normalization may differ from raw Outcome_Data by design.
    # Only changes relative to v33 source are propagated, not wholesale overwrites.
    if rid in byid:
        cur=byid[rid];orig=next(t for t in old if t['Comparison ID']==rid)
        for h in OH:
            if cur[h]!=orig[h]:z[h]=cur[h]
        for h in EXTRA:z[h]=cur.get(h)
    else:z['V34 record ID']='AF-HIST-'+rid;z['V34 eligibility']='HISTORICAL_AF_ONLY';z['V34 eligibility reason']='Retained source-normalized AF result; original result-specific lock applies, no automatic new pooling.'
    af.append(z)
for r in outcomes:
    if r['Comparison ID'] not in afids:af.append(dict(r))
updates['Outcome_Data_AF_LOCK']=matrix(af,v['Outcome_Data_AF_LOCK'][0]+EXTRA)
afby={r['Comparison ID']:r for r in af};sl=[]
for r in records(v['Stata_AF_Long']):
    z=dict(r);cur=afby.get(r['comparison_id'])
    if cur:
        for k,h in native_map.items():
            if k in z:z[k]=cur.get(h)
        z['study']=cur['Canonical study'];z['time_window']=cur['Timepoint/window'];z['v34_record_id']=cur.get('V34 record ID')
        if cur.get('V34 eligibility') in ['HOLD','EXCLUDE','GRAPH ONLY','NOT REPORTED','SOURCE NOT ACCESSED','DUPLICATE HOLD']:
            z.update(include_strict=0,include_sensitivity=0,narrative_only=1,analysis_action='V34_HOLD',hold_reason=cur.get('V34 eligibility reason'))
    sl.append(z)
for r in outcomes:
    if r.get('V34 audit row') or r.get('V34 disposition')=='SOURCE_NOT_ACCESSED':
        z={h:None for h in v['Stata_AF_Long'][0]};z.update({k:r.get(h) for k,h in native_map.items() if k in z})
        z.update(lock_id=None,target='V34_NEW_RESULT',endpoint_stratum=r['V34 endpoint class'],study=r['Canonical study'],time_window=r['Timepoint/window'],reported_p=r['Reported P'],rob_overall='ROB2_RESULT_SPECIFIC_PENDING',include_strict=0,include_sensitivity=0,narrative_only=1,analysis_action='See V34_Secondary_Ready; no inherited AF lock',hold_reason=r['V34 eligibility reason'],source_qc=r['Source-QC issue'],record_status=r['Record status'],v34_record_id=r['V34 record ID']);sl.append(z)
updates['Stata_AF_Long']=matrix(sl,v['Stata_AF_Long'][0]+['v34_record_id'])
writecsv(BASE/'data/v34_stata_af_long.csv',sl,v['Stata_AF_Long'][0]+['v34_record_id'])

# Remaining workbook views and explicit source audit layers.
studycounts=Counter(r['Canonical study'] for r in outcomes)
sm=records(v['Study_Master'])
for r in sm:
    r['Source-normalized outcome rows']=studycounts[r['Canonical study']]
    if r['Canonical study'] in src:
        r['Source-summary status']='V34 SOURCE-AUDIT RECONCILED; result-level holds visible'
        r['Sol/source correction note']=str(r['Sol/source correction note'] or '')+' | V34: '+src[r['Canonical study']]['verification_notes']
updates['Study_Master']=matrix(sm,v['Study_Master'][0])
ct=Counter(r['V34 eligibility'] for r in outcomes)
family=defaultdict(list)
for r in outcomes:family[r['V34 endpoint class']].append(r)
updates['Analysis_Summary']=matrix([{'Synthesis family':f,'Total rows':len(rs),'Include-ready':sum(r['V34 eligibility']=='INCLUDE' for r in rs),'Conditional/secondary':sum(r['V34 eligibility'] in ['CONDITIONAL','NARRATIVE ONLY'] for r in rs),'Excluded/data unavailable':sum(r['V34 eligibility'] not in ['INCLUDE','CONDITIONAL','NARRATIVE ONLY'] for r in rs),'Key note':'Native candidates; separate exact endpoint, window, modality, comparator, statistic and shared arms.','Dedicated candidate set':'V34_Secondary_Ready','Rows':sum(r['V34 eligibility']=='INCLUDE' for r in rs)} for f,rs in sorted(family.items())],v['Analysis_Summary'][0])
updates['Study_Level_Synthesis']=matrix([dict(zip(v['Study_Level_Synthesis'][0],[f,len(rs),len({r['Canonical study'] for r in rs}),sum(r['V34 eligibility']=='INCLUDE' for r in rs),sum(r['V34 eligibility'] in ['HOLD','CONDITIONAL'] for r in rs),sum(r['V34 eligibility']=='DUPLICATE HOLD' for r in rs),'Candidate rows are not independent studies; native estimands retained.'])) for f,rs in sorted(family.items())],v['Study_Level_Synthesis'][0])
summary=[list(r) for r in v['Summary']]
summary[0][0]='Reconciled Master v34 — QC and Readiness'
metrics=[('Canonical studies/reports',len(CANON)),('Outcome_Data rows',len(outcomes)),('v33 outcome rows',len(old)),('Net rows added',len(outcomes)-len(old)),('Audited studies',len(src)),('Audit rows dispositioned',len(dispositions)),('Strict primary k',len(primary)),('Analysis-extractable native rows',ct['INCLUDE']),('Conditional rows',ct['CONDITIONAL']),('Narrative-only rows',ct['NARRATIVE ONLY']),('Graph-only rows',ct['GRAPH ONLY']),('Not-reported rows',ct['NOT REPORTED']),('Source-not-accessed outcome rows',ct['SOURCE NOT ACCESSED']),('Unresolved / QC hold rows',ct['HOLD']),('Duplicate-family hold rows',ct['DUPLICATE HOLD']),('Excluded rows',ct['EXCLUDE'])]
summary=summary[:3]+[[k,n,None,None,None,None,None,None] for k,n in metrics]
summary[2]=['Metric','Count',None,'Status',None,None,None,None]
summary[3][3]='Source audit reconciled; unresolved results remain on hold.'
updates['Summary']=summary
readme=[list(r) for r in v['README']]
readme[0][0]='TEAS/EA Systematic Review — Reconciled Master Data v34'
for row in readme:
    if row[0]=='Version':row[1]='v34 — 2026-09-08'
    if row[0]=='Purpose':row[1]='Reconcile 374 audit records into frozen v33; full source quotations, row dispositions, raw statistics and visible source conflicts. Supplied source-PDF audit used; PDFs not independently re-extracted in this run.'
    if row[0]=='Final dataset status':row[1]='RECONCILIATION READY WITH EXPLICIT HOLDS. Not all secondary outcomes are poolable or RoB-assessed. See V34_QC and V34_Source_Conflicts.'
    if row[0]=='Analysis readiness':row[1]='V34 INCLUDE means a numeric native result is extractable, not permission to pool across endpoints/windows/arms. V34_Secondary_Ready records explicit selections. New eligible results have result-specific RoB pending.'
    if row[0]=='Data preservation':row[1]='v33 remains byte-identical. Historical AF_Result_Lock and related disposition records are retained; original AF data and Stata long snapshots are carried in V33_* sheets. Active corrections are fully logged.'
    if row[0]=='A–F reconciliation':row[1]='Historical A–F adjudications retained. New results receive no inherited AF judgment; active source corrections propagate to regenerated AF/Stata layers.'
    if row[0]=='Stata status':row[1]='Strict primary regenerated from seven locked source rows; native secondary sets regenerated without median conversion, graph estimation, percentage back-calculation, or proxy opioid MME.'
readme += [[k,val]+[None]*6 for k,val in [('V34 BASE SHA256',vi['sha256']),('V34 AUDIT SHA256',ai['sha256']),('V34 OUTCOME ROWS',len(outcomes)),('V34 source status rule','NOT REPORTED and SOURCE NOT ACCESSED remain distinct. Generic unaccessed protocol/supplement statuses live in V34_Source_PDF_Audit; only identified missing endpoints add Outcome_Data rows.'),('V34 history rule','AF_Result_Lock, AF_P1_Disposition, Corrected_RoB2 and Scope_Exclusion_Audit retain original adjudications. Active status overrides are in Stata_AF_Long, V34_Result_RoB2 and V34_Source_Conflicts.'),('V34 output folder',str(BASE))]]
updates['README']=readme
crit=records(a['Critical_QC'])
sources=[]
for st,r in src.items():sources.append(dict(r,critical_qc=' | '.join(x['issue']+' HANDLING: '+x['recommended_handling'] for x in crit if x['study']==st),verification_basis='Supplied consolidated audit. No independent PDF reread in this reconciliation run.'))
updates['V34_Source_PDF_Audit']=matrix(sources)
updates['V34_Source_Conflicts']=matrix(conf,CH)
audit_full=[]
for x,d in zip(aud,dispositions):audit_full.append(dict(x,**{f'v34_{k}':val for k,val in d.items() if k not in ['batch','batch_row','study']}))
updates['V34_Audit_Dispositions']=matrix(audit_full)
updates['V34_Reconciliation_Status']=a['Reconciliation_Status']
updates['V34_Composite_Matching']=matrix(comparisons)
updates['V34_Result_RoB2']=matrix([{'record_id':r['V34 record ID'],'study':r['Canonical study'],'comparison_id':r['Comparison ID'],'outcome':r['Outcome/result'],'window':r['Timepoint/window'],'eligibility':r['V34 eligibility'],'result_specific_rob2':r['V34 RoB2 status']} for r in outcomes])
writecsv(BASE/'data/v34_audit_dispositions.csv',audit_full)
writecsv(BASE/'data/v34_outcome_data.csv',outcomes,OH+EXTRA)
writecsv(BASE/'data/v34_composite_matching.csv',comparisons)
writecsv(BASE/'data/v34_source_conflicts.csv',conf,CH)

# Save stage for analysis preparation and workbook authoring.
state={'updates':updates,'outcomes':outcomes,'primary':primary,'dispositions':dispositions,'conflicts':conf,'correction_reason':correction_reason,'correction_type':correction_type,'counts':dict(ct),'duplicate_keys_existing':dups,'inputs':{'v33':vi,'audit':ai}}
(BASE/'data/reconciliation_stage.pending.json').write_text(json.dumps(state,ensure_ascii=False,indent=2))
(BASE/'data/reconciliation_stage.pending.json').replace(BASE/'data/reconciliation_stage.json')
print(j({'v33':len(old),'v34':len(outcomes),'dispositions':dict(Counter(d['disposition'] for d in dispositions)),'readiness':dict(ct),'conflicts':len(conf),'primary_k':len(primary),'introduced_exact_duplicate_keys':introduced}))
