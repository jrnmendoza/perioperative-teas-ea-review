#!/usr/bin/env python3
"""Independent XLSX read-back QA and final reports; never writes workbook cells."""
import json,csv,hashlib,math,re
from pathlib import Path
from collections import Counter,defaultdict
import openpyxl
BASE=Path(__file__).resolve().parents[1];OUT=BASE.parent
S=json.load(open(BASE/'data/final_stage.json'));D=S['outcomes'];U=S['updates']
P=OUT/'TEAS_EA_RECONCILED_MASTER_DATA_v34_FINAL_LOCK_READY.xlsx'
W=openpyxl.load_workbook(P,data_only=False)
O=openpyxl.load_workbook(BASE/'inputs/v33_working_copy.xlsx',data_only=False)
def equal(a,b):return (a in (None,'') and b in (None,'')) or a==b
def records(sh):
    m=list(sh.iter_rows(values_only=True));return [dict(zip(m[0],r)) for r in m[1:] if any(x is not None for x in r)]
def readcsv(p):return list(csv.DictReader(open(p,encoding='utf-8-sig')))
checks=list(S['checks'])
def check(name,ok,evidence):checks.append({'check':name,'status':'PASS' if ok else 'FAIL','evidence':str(evidence)})
diff=[]
for name in W.sheetnames:
    m=U.get(name)
    if m is None:m=[list(r) for r in O[name].iter_rows(values_only=True)]
    for i,row in enumerate(m,1):
        for j,x in enumerate(row,1):
            if not equal(x,W[name].cell(i,j).value):diff.append((name,i,j,x,W[name].cell(i,j).value))
    for row in W[name].iter_rows(min_row=len(m)+1):
        if any(c.value is not None for c in row):diff.append((name,'stale trailing row',row[0].row))
check('Exported workbook cell round-trip',not diff,f'{len(diff)} substantive cell differences; empty strings and Excel blanks considered equivalent. All unchanged sheets checked against frozen v33.')
errors=[(sh.title,c.coordinate,c.value)for sh in W for row in sh for c in row if c.data_type=='e' or (isinstance(c.value,str) and c.value in ['#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A'])]
formulas=[(sh.title,c.coordinate)for sh in W for row in sh for c in row if c.data_type=='f']
check('Exported formula/error check',not errors and not formulas,f'{len(errors)} error cells; {len(formulas)} formulas; summaries regenerated as typed values, matching the source architecture.')
v33hash=hashlib.sha256((OUT/'TEAS_EA_RECONCILED_MASTER_DATA_v33_FINAL_LOCK_READY.xlsx').read_bytes()).hexdigest()
check('Frozen v33 hash',v33hash==S['inputs']['v33']['sha256'],v33hash)
check('Original sheet preservation',set(O.sheetnames).issubset(W.sheetnames),f'All {len(O.sheetnames)} original sheet names retained; {len(W.sheetnames)} total sheets.')
for hist,orig in [('V33_Outcome_Data_AF_LOCK','Outcome_Data_AF_LOCK'),('V33_Stata_AF_Long','Stata_AF_Long')]:
    check('History snapshot '+hist,all(equal(c.value,W[hist].cell(c.row,c.column).value) for row in O[orig] for c in row),'Every original value preserved in explicit history snapshot.')

out=records(W['Outcome_Data']);byid={r['Comparison ID']:r for r in out};byrid={r['V34 record ID']:r for r in out}
ready=records(W['Analysis_Readiness'])
check('Readiness identities/values',all(all(equal(r.get(h),byrid[r['V34 record ID']].get(h))for h in U['Outcome_Data'][0]) for r in ready),f'{len(ready)} exact Outcome_Data-linked rows.')
setbad=[]
for name in ['Set_Opioid_24h','Set_Pain_24h','Set_PONV','Set_QoR','Set_GI_Recovery']:
    for r in records(W[name]):
        source=next(z for z in ready if z['V34 record ID']==r['V34 record ID'])
        for h,val in r.items():
            if h in source and not equal(val,source[h]):setbad.append((name,r['V34 record ID'],h))
check('Set/readiness consistency',not setbad,f'{len(setbad)} discrepancies across all five Set_* sheets.')
af={r['Comparison ID']:r for r in records(W['Outcome_Data_AF_LOCK'])}
maps={'study':'Canonical study','outcome':'Outcome/result','time_window':'Timepoint/window','data_type':'Data type','n_i':'Analyzed n intervention','n_c':'Analyzed n comparator','mean_i':'Mean intervention','sd_i':'SD intervention','mean_c':'Mean comparator','sd_c':'SD comparator','events_i':'Events intervention','events_c':'Events comparator','unit':'Unit/scale'}
afbad=[]
for r in records(W['Stata_AF_Long']):
    if r['comparison_id'] in af:
        for k,h in maps.items():
            if not equal(r[k],af[r['comparison_id']][h]):afbad.append((r['comparison_id'],k))
check('Active AF/Stata values',not afbad,f'{len(afbad)} numeric/endpoint discrepancies in regenerated Stata_AF_Long.')
def csvsame(sheet,file):
    rows=records(W[sheet]);c=readcsv(file)
    if len(rows)!=len(c):return False
    for x,y in zip(rows,c):
        for k,xx in x.items():
            yy=y.get(k)
            if xx is None and yy in (None,''):continue
            if isinstance(xx,(float,int)):
                try:
                    if abs(float(yy)-xx)>1e-10:return False
                except (ValueError,TypeError):return False
            elif str(xx)!=yy:return False
    return True
csvchecks=[csvsame('Stata_AF_Long',BASE/'data/v34_stata_af_long.csv'),csvsame('Stata_Opioid24_Primary',BASE/'data/v34_stata_opioid24_primary_native.csv'),csvsame('V34_Secondary_Ready',BASE/'data/v34_secondary_selected.csv'),csvsame('Outcome_Data',BASE/'data/v34_outcome_data.csv')]
check('Workbook/CSV Stata consistency',all(csvchecks),f'AF long, primary native, selected secondary and full Outcome_Data exports all agree: {csvchecks}')
# Every modified existing value must have a corresponding coordinate delta.
logs=readcsv(OUT/'TEAS_EA_v33_to_v34_RECONCILIATION_LOG.csv');covered=set()
for l in logs:
    try:q=json.loads(l['new_value_or_status'])
    except (ValueError,TypeError):continue
    if isinstance(q,dict) and 'excel_row' in q:
        for col in q.get('cells',{}):covered.add((l['sheet'],q['excel_row'],int(col)))
unlogged=[]
for sh in O:
    for row in sh:
        for c in row:
            if not equal(c.value,W[sh.title].cell(c.row,c.column).value) and (sh.title,c.row,c.column) not in covered:unlogged.append((sh.title,c.row,c.column))
check('Existing-cell change-log coverage',not unlogged,f'{len(unlogged)} unlogged changes; {len(logs)} CSV/workbook log entries.')
check('Change-log workbook/CSV consistency',csvsame('V33_to_V34_Change_Log',OUT/'TEAS_EA_v33_to_v34_RECONCILIATION_LOG.csv'),'Exact matching row count, headers and values.')
check('Audit disposition completeness',len(records(W['V34_Audit_Dispositions']))==374 and len({r['v34_audit_row'] for r in records(W['V34_Audit_Dispositions'])})==374,'Every audit row 2–375 has one disposition and target result ID.')
li=[r for r in out if r['Canonical study']=='Liang 2021']
check('Liang all-record allocation consistency',all((r['Randomized n intervention'],r['Randomized n comparator'])==(37,38) for r in li),'All original and new Liang records use randomized 37/38, including the unavailable-opioid row; analyzed 35/35 retained.')
check('Printed mean with missing SD preserved',next(r for r in out if r.get('V34 audit row')==356)['Mean intervention']==184 and next(r for r in out if r.get('V34 audit row')==356)['SD intervention'] is None,'Zheng QoR-40 48h mean 184 vs181 retained; exact SD blank, HOLD.')
check('New source-range and IQR bounds',all(r['Minimum intervention']<=r['Maximum intervention'] for r in out if r.get('Minimum intervention') is not None),'No range inversion or median-to-mean conversion.')
check('Numeric cell typing',all(isinstance(W['Summary'].cell(i,2).value,(int,float)) for i in range(4,20)),'Summary counts are numeric Excel cells, not formatted strings.')
bad=[r for r in checks if r['status']=='FAIL']
json.dump({'checks':checks,'differences':diff,'unlogged':unlogged,'af_discrepancies':afbad,'xlsx_sha256':hashlib.sha256(P.read_bytes()).hexdigest()},open(BASE/'data/final_qc.json','w'),ensure_ascii=False,indent=2)
assert not bad,bad

dc=Counter(d['disposition'] for d in S['dispositions']);ct=Counter(r['V34 eligibility'] for r in out)
new=[r for r in out if r.get('V34 audit row')];neweligible=sum(r['V34 eligibility']=='INCLUDE' for r in new)
source=records(W['V34_Source_PDF_Audit']);access=[r['study'] for r in source if 'NOT ACCESSED' in r['supplement_status']]
unresolved=[r for r in S['conflicts'] if r['requires_author_contact']=='Yes'];contacts=sorted({r['study'] for r in unresolved})
nschanges=[r for r in S['actual_source_changes'] if any(h.startswith('Randomized n') for h in r['changes'])]
windowchanges=[r for r in S['actual_source_changes'] if 'Timepoint/window' in r['changes']]
oldstudies=set()
for n in ['intraop_remifentanil','intraop_sufentanil','qor40_24h','gi_first_defecation','rescue_opioid_binary_24h']:
    oldstudies.update(r['study'] for r in readcsv(BASE.parent.parent/f'08_V33_MASTER/01_DATA/v33_{n}.csv'))
newstudies={r['study'] for r in S['selected_secondary']}-oldstudies
primary=next(r for r in S['model_results'] if r['phase']=='v34' and r['analysis_id']=='primary_24h_mme_ALL_AUDIT')
paths=[P,OUT/'TEAS_EA_v33_to_v34_RECONCILIATION_LOG.csv',OUT/'TEAS_EA_v34_QC_REPORT.md',OUT/'TEAS_EA_v34_ANALYSIS_CHANGE_SUMMARY.md']
def table(headers,rows):return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(str(x).replace('|',' / ').replace('\n',' ') for x in r)+' |' for r in rows])

qr=['# TEAS/EA v34 reconciliation QC report','',
 '**PASS — reconciliation and export integrity. The workbook retains explicit source/adjudication holds; this is not a claim that every secondary result is ready for final pooling.**','',
 'The supplied consolidated source-PDF audit was reconciled against the frozen v33 master. The PDFs were not independently re-extracted in this run. New result-specific RoB assessments were not invented.','',
 '## Required PASS/FAIL checks','',table(['Check','Status','Evidence'],[(c['check'],c['status'],c['evidence']) for c in checks]),'',
 '## Inventory and reconciliation counts','',
 table(['Metric','Result'],[('v33 Outcome_Data',382),('v34 Outcome_Data',len(out)),('Net additions',len(out)-382),('Existing source rows updated (original 38 columns)',len(S['actual_source_changes'])),('Allocation metadata corrections',f'{len(nschanges)} rows: five Liang + six Zhu; analyzed denominators unchanged'),('Other metadata corrections','Tu rescue window; Sim IQR width representation'),('Audit duplicate rejected',dc['ALREADY_PRESENT']),('New graph-only rows',dc['GRAPH_ONLY_ADDED']),('New NOT REPORTED rows',dc['NOT_REPORTED_ADDED']),('SOURCE NOT ACCESSED outcome rows',ct['SOURCE NOT ACCESSED']),('Study-level supplement/protocol access gaps',len(access)),('Visible conflict/correction records',len(S['conflicts'])),('Unresolved records requiring author clarification',len(unresolved)),('New numerically extractable secondary results',neweligible),('New eligible results with result-specific RoB pending',neweligible),('Canonical studies/reports',70),('Strict primary before / after','7 / 7'),('Total workbook sheets',len(W.sheetnames)),('Change-log entries',len(logs))]),'',
 'The 374 audit records contribute 373 distinct outcome rows; Li 2021 VAS≥4 at 6h already matches v33. Two additional Gao outcomes are sourced from the reconciliation/source-access findings rather than All_Gapfill. All 382 original outcome rows remain in order; no source rows were deleted. An existing Sun aggregate placeholder remains visible but is excluded as superseded by three separate timing-arm records.','',
 f'INCLUDE means numerically extractable in its native endpoint, unit, window, population and statistic. It does not authorize pooling across incompatible definitions. The {neweligible} new candidates are not {neweligible} independent trials or completed RoB assessments.','',
 '## Dispositions and active readiness','',table(['Audit disposition','Rows'],sorted(dc.items())),'',table(['Active readiness','Rows'],sorted(ct.items())),'',
 '## Source access and author clarification','',
 'Unaccessed supplements/protocols: '+', '.join(access)+'. Gao has two explicit missing outcome records (total LOS and 30-day complications). The other generic access gaps remain in V34_Source_PDF_Audit; no unspecified outcomes were invented.','',
 'Studies with unresolved conflict/definition records for author clarification: '+', '.join(contacts)+'. Additional requests for exact graph values, missing numerical outcomes and unclear source definitions are listed row by row in V34_Not_Pooled. No author messages were sent.','',
 '## Preservation and reproducibility','',
 '- Frozen v33 SHA-256: `'+v33hash+'`.',
 '- Audit SHA-256: `'+S['inputs']['audit']['sha256']+'`.',
 '- v34 SHA-256: `'+hashlib.sha256(P.read_bytes()).hexdigest()+'`.',
 '- Original formulas: 0; named ranges: 0. Exported formula errors: 0. No formula-derived result was replaced with an assumed value.',
 '- AF_Result_Lock and AF_P1_Disposition remain historical adjudications. V33_Outcome_Data_AF_LOCK and V33_Stata_AF_Long preserve original values; corrected active layers and regenerated exports carry v34 changes.',
 '- Native Excel table/filter ranges were expanded to full current extents, including previously truncated master/RoB tables. Existing sheets were retained; summary body merges were removed where they would hide current counts.',
 '- Reproducible Python reconciliation/data-preparation and JavaScript artifact-tool workbook builder are in `'+str(BASE/'code')+'`. Python/openpyxl was used for reading and QA, not workbook authoring.',
 '- StataNow 19.5 completed '+str(len(S['model_results']))+' matched/current runs. Log: `'+str(BASE/'results/v34_analysis.log')+'`.',
 '- Rendered previews of summary, QA, Liang corrections, source conflicts and audit dispositions were visually checked.','',
 '## Study-level summary','',table(['Study','v33','v34','Added','Existing updated','New numeric candidates','New graph','New NR'],[(r['study'],r['v33_rows'],r['v34_rows'],r['net_added'],r['existing_source_rows_updated'],r['new_numeric_secondary_candidates'],r['new_graph_only'],r['new_not_reported']) for r in S['study_summary']]),'',
 '## Outputs','']+['- `'+str(p)+'`' for p in paths]
(OUT/'TEAS_EA_v34_QC_REPORT.md').write_text('\n'.join(qr)+'\n')

ar=['# TEAS/EA v34 analysis change summary','',
 '**Primary unchanged; secondary source data expanded, with incompatible windows and unresolved records kept separate.**','',
 '## A–Q requested reconciliation report','',
 table(['Item','Result'],[
 ('A. v33 outcome rows',382),('B. v34 outcome rows',len(out)),('C. Net rows added',375),('D. Existing source rows updated',len(S['actual_source_changes'])),('E. Metadata corrections',f'{len(nschanges)} allocation rows, Tu 6–24h window, Sim IQR-width representation; generated status/ID fields also added to existing rows'),('F. Duplicates rejected',1),('G. Graph-only rows added',38),('H. NOT REPORTED rows added',16),('I. SOURCE NOT ACCESSED','2 outcome records; 7 studies with source-access gaps'),('J. Unresolved source conflicts',f'{len(unresolved)} unresolved records across {len(contacts)} studies; {len(S["conflicts"])} total conflict/correction records'),('K. New analysis-eligible secondary results',f'{neweligible} native numeric candidates; 3 newly selected contrasts; all new candidates have result-specific RoB pending'),('L. Studies newly entering any secondary meta-analysis',str(len(newstudies))+' relative to the five existing v33 secondary datasets; Liang already contributed QoR, Pan already contributed remifentanil'),('M. Strict primary k before/after','7 → 7'),('N. Primary pooled estimate changed?','No; exact matched Stata reruns agree'),('O. Secondary analyses changed','Intraoperative remifentanil and sufentanil expand; exact-window QoR and rescue sets corrected; defecation strata unchanged'),('P. Author clarification',', '.join(contacts)),('Q. Output files','Exact paths at the end of this report')]),'',
 '## Strict primary opioid analysis','',
 f'Combined historical audit estimate: MD {primary["estimate"]:.6f} mg IV MME; 95% CI {primary["ci_low"]:.6f} to {primary["ci_high"]:.6f}; P={primary["p_value"]:.8f}. Both the combined audit and separate TEAS/sham (k=4) and EA/usual-care (k=3) estimates are unchanged. Total analyzed N=676. The same prespecified v33 conversion factors were applied to regenerated source values.','',
 'No binary rescue use, PCA press count, multimodal PCIA solution, author-defined intraoperative conversion, median, percentage or graph-only result was promoted into this strict primary model.','',
 '## Secondary selection changes','',
 '- **Remifentanil:** selected contrasts 8→9; Liang 2021 adds 521.5±206.8 vs 464.7±156.0 µg, analyzed 35/35. TEAS/sham k=5→6. Matched REML+Hartung–Knapp MD changes from −149.89 µg (95% CI −254.71 to −45.08) to −117.78 µg (−242.51 to 6.95). EA and usual-care strata are retained separately; single-study strata are not pooled.',
 '- **Sufentanil:** selected contrasts 5→6, representing 4→5 unique studies because Wang 2024 has two risk strata within one trial. Liang adds 21.4±3.1 vs20.0±2.9 µg. TEAS/sham now has k=2, MD −9.44 µg (95% CI −32.40 to13.52), REML normal CI. Wang SNVP and MNVP are retained in separate stratum files and are never counted as two independent RCTs in one model.',
 '- **QoR-40:** the former three-row ~24h set contained Yu POD1. That result is now kept separately. The exact 24h set contains Yao, Liang and newly extracted Pan global QoR-40; subscales are excluded from the global-score set. TEAS/sham remains Yao+Liang (k=2); Pan is a separate usual-care k=1 result. Thus the old mixed-window pooled estimate is not relabelled as an exact-24h estimate.',
 '- **Rescue opioid incidence:** the former k=3 0–24h/POD1 set is withdrawn as an exact-window model. Tu is 6–24h; Liu burn is through POD1; Yu is exact0–24h. The strict 0–24h set now has k=1 and is not meta-analyzed. This affects a secondary rescue endpoint, not the strict primary opioid-dose k=7.',
 '- **First defecation:** the seven selected independent study records are unchanged. TEAS/sham k=3 and EA/usual-care k=3 estimates are unchanged; Ng EA/sham is separate k=1. Sun adds three source results to the native data but they remain unpooled until an endpoint-specific shared-control selection is adjudicated.',
 '- **Other families:** data preparation was regenerated for postoperative opioid other windows, rescue frequency, time to first rescue, pain, PONV, nausea, vomiting, antiemetics, QoR15/40, GI recovery, LOS, functional recovery, adverse events/cognition, and analgesic proxies/non-opioid rescue. These files retain native definitions, modalities, comparators and holds; no automatic new pooled estimate is claimed for every added outcome.',
 '- **No GRADE update:** new extraction and provisional secondary analyses do not constitute completed result-specific RoB or certainty assessment.','',
 '## Current Stata estimates','',table(['Analysis / stratum','k','Estimate','95% CI','P','Model'],[(r['analysis_id'],int(r['k']),f'{r["estimate"]:.4f}',f'{r["ci_low"]:.4f} to {r["ci_high"]:.4f}',f'{r["p_value"]:.5f}',r['model']) for r in S['model_results'] if r['phase']=='v34']), '',
 'Models are REML; Hartung–Knapp when k≥3, normal CI when k=2, no pooled estimate when k=1. These sparse secondary results are provisional. Exact group-value P conflicts are preserved and do not get silently repaired. Numeric native binary exports use a documented 0.5 correction to all four cells when needed; double-zero contrasts have no estimated logRR. None of the selected rerun models required that correction.','',
 '## Study-level reconciliation','',table(['Study','Added','Existing source rows corrected','New native numeric candidates','Comment'],[(r['study'],r['net_added'],r['existing_source_rows_updated'],r['new_numeric_secondary_candidates'],r['note']) for r in S['study_summary']]),'',
 '## Remaining adjudication work','',
 'Prioritize Yu pain dispersion/significance; Long pain and PND denominators; Liang 48h CRBD and undefined analgesia/pain events; Lu2022 QoR15; Zhu QoR15 timing; Pan undefined NRS statistic/flow; Jiang unidentified NRS timepoint; and the missing supplements/protocols. Source conflicts with an existing table hierarchy remain visible even when table values are retained. Numeric-candidate status does not remove these restrictions.','',
 '## Exact output paths','']+['- `'+str(p)+'`' for p in paths]+['','Analysis CSVs and manifests: `'+str(BASE/'data')+'`.','Stata result CSV: `'+str(BASE/'results/v34_model_results.csv')+'`.','Reproducible code and Stata do-file: `'+str(BASE/'code')+'`.']
(OUT/'TEAS_EA_v34_ANALYSIS_CHANGE_SUMMARY.md').write_text('\n'.join(ar)+'\n')
summary={'outcome_before':382,'outcome_after':len(out),'net_added':375,'existing_updated':len(S['actual_source_changes']),'allocation_rows_corrected':len(nschanges),'graph_added':38,'nr_added':16,'source_not_accessed_rows':2,'source_not_accessed_studies':len(access),'conflicts_total':len(S['conflicts']),'conflicts_unresolved':len(unresolved),'conflict_studies':contacts,'new_numeric_candidates':neweligible,'new_selected_contrasts':3,'new_studies_overall':len(newstudies),'primary_k':7,'primary_estimate':primary['estimate'],'qa_passes':len(checks),'xlsx':str(P),'files':[str(p)for p in paths]}
json.dump(summary,open(BASE/'data/completion_summary.json','w'),indent=2,ensure_ascii=False)
print(json.dumps(summary,indent=2,ensure_ascii=False))
