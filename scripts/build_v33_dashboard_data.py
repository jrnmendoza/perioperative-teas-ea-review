#!/usr/bin/env python3
"""
Generate dashboard/v33_data.js -- the v33 machine-readable analytical layer.

Two things live here:

1. STUDY CONTRIBUTION MAP. Which of the 70 included RCTs contribute to which
   outcome family, derived from v33 Outcome_Data. Its job is to make plain why
   "70 included RCTs" does not mean "70 studies in the primary meta-analysis":
   a trial can be fully included in the review and still carry no analysable
   0-24 h cumulative opioid dose.

2. v33 SECONDARY RESULTS, read from the Stata output. No number is copied from
   old HTML or JS.

Every count is derived. Nothing about the map is hardcoded except the display
order of the outcome families.
"""

from __future__ import annotations

import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
V33 = (ROOT / "TEAS EA Verification"
       / "TEAS_EA_RECONCILED_MASTER_DATA_v33_FINAL_LOCK_READY.xlsx")
SEC = ROOT / "08_V33_MASTER" / "03_RESULTS" / "results_v33_secondary.csv"
NOTPOOLED = ROOT / "08_V33_MASTER" / "01_DATA" / "v33_not_pooled_register.csv"
PRIMARY = ROOT / "06_FINAL_ANALYSIS_V26" / "03_RESULTS" / "results_opioid24_primary.csv"
OPIOID_CSV = ROOT / "06_FINAL_ANALYSIS_V26" / "01_DATA" / "opioid_24h_primary.csv"
OUT = ROOT / "dashboard" / "v33_data.js"
FIGURES = {
    'V33_RESCUE_OPIOID_RR_24H': 'forest_v33_rescue_opioid_rr.png',
    'V33_INTRAOP_REMI_MD': 'forest_v33_intraop_remifentanil.png',
    'V33_INTRAOP_SUF_MD': 'forest_v33_intraop_sufentanil.png',
    'V33_QOR40_24H_MD': 'forest_v33_qor40_24h.png',
    'V33_GI_DEFECATION_MD': 'forest_v33_gi_defecation.png',
}

# Display order and human labels for the contribution map.
FAMILY_GROUPS = [
    ("primary_opioid_24h", "Primary 0–24 h opioid",
     "Contributes an analysable cumulative 0–24 h postoperative opioid dose to the strict primary model."),
    ("other_postop_opioid", "Other postoperative opioid",
     "Reports postoperative opioid consumption at another window, in another unit, or in a form that cannot enter the strict model."),
    ("intraoperative_opioid", "Intraoperative opioid",
     "Reports intraoperative opioid requirement. A different estimand from postoperative consumption; never pooled with it."),
    ("rescue_opioid", "Rescue analgesia / opioid",
     "Reports rescue opioid use, rescue administration counts, or time to first rescue."),
    ("opioid_demand", "PCA demand / presses",
     "Reports PCA attempts, presses or deliveries. A behavioural proxy, not a drug quantity."),
    ("pain", "Pain", "Reports a postoperative pain score."),
    ("ponv", "PONV / nausea / vomiting", "Reports postoperative nausea or vomiting."),
    ("qor", "Quality of recovery", "Reports QoR-40 or QoR-15."),
    ("gi_recovery", "GI recovery", "Reports flatus, defecation, bowel sounds, diet or ileus."),
    ("los_recovery", "LOS / recovery", "Reports length of stay or another recovery endpoint."),
    ("other_narrative", "Other / narrative only",
     "Carries outcome information that no current pooled model uses."),
]

FAMILY_MAP = {
    "Opioid consumption": "other_postop_opioid",
    "Postoperative opioid consumption": "other_postop_opioid",
    "Perioperative opioid": "other_postop_opioid",
    "Intraoperative opioid": "intraoperative_opioid",
    "Intraoperative anesthetic": "intraoperative_opioid",
    "Rescue opioid use": "rescue_opioid",
    "Rescue analgesia": "rescue_opioid",
    "Opioid demand": "opioid_demand",
    "PCA demand/proxy": "opioid_demand",
    "Pain": "pain",
    "Experimental pain": "pain",
    "Chronic pain": "pain",
    "PONV": "ponv",
    "Rescue antiemetic": "ponv",
    "Quality of recovery": "qor",
    "GI recovery": "gi_recovery",
    "Recovery": "los_recovery",
    "Functional recovery": "los_recovery",
    "Urinary recovery": "los_recovery",
    "Sleep quality": "other_narrative",
    "Neurocognitive": "other_narrative",
    "Catheter-related bladder discomfort": "other_narrative",
    "Other adverse events": "other_narrative",
    "Cardiac rhythm": "other_narrative",
}


def read_csv(p: Path) -> list[dict]:
    with p.open(encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def has_data(r: dict) -> bool:
    for k in ("Mean intervention", "Median intervention", "Events intervention"):
        if r.get(k) not in (None, ""):
            return True
    return False


def main() -> int:
    wb = openpyxl.load_workbook(V33, data_only=True)

    sm = list(wb["Study_Master"].iter_rows(values_only=True))
    sm_hdr = [str(h) for h in sm[0]]
    studies = [dict(zip(sm_hdr, r))["Canonical study"] for r in sm[1:]]
    studies = [s for s in studies if s]

    od = list(wb["Outcome_Data"].iter_rows(values_only=True))
    od_hdr = [str(h) for h in od[0]]
    rows = [dict(zip(od_hdr, r)) for r in od[1:]]

    # strict primary contributors, straight from the locked analysis dataset
    strict = {r["study_unit"] for r in read_csv(OPIOID_CSV) if r.get("inc_primary") == "1"}

    contrib: dict[str, set] = defaultdict(set)
    for r in rows:
        s = r.get("Canonical study")
        if not s or not has_data(r):
            continue
        g = FAMILY_MAP.get(str(r.get("Outcome family")), "other_narrative")
        contrib[s].add(g)

    # A study is a primary contributor only if the locked dataset names it
    # EXACTLY. Prefix matching is wrong here and was caught doing real damage:
    # "He 2026 (breast/WJCO)" and "He 2026 (hepatectomy/JIS)" are different
    # trials, and a stem match on "He 2026" credited the breast trial with the
    # hepatectomy trial's primary contribution, inflating the map to 8 against
    # a locked k of 7.
    for s in list(contrib):
        if s in strict:
            contrib[s].add("primary_opioid_24h")

    # A trial whose only rows are graph-only or "not reported" carries no
    # analysable number anywhere. It is still an included RCT and must stay
    # visible rather than silently dropping out of the map -- it is the
    # clearest illustration of why 70 included does not mean 70 analysable.
    graph_only = []
    for s in studies:
        if s in contrib:
            continue
        if any(r.get("Canonical study") == s for r in rows):
            contrib[s].add("other_narrative")
            graph_only.append(s)

    per_study = []
    for s in sorted(studies):
        fams = sorted(contrib.get(s, set()))
        per_study.append({
            "study": s,
            "families": fams,
            "n_families": len(fams),
            "in_primary": "primary_opioid_24h" in fams,
        })

    group_counts = {gid: sum(1 for x in per_study if gid in x["families"])
                    for gid, _, _ in FAMILY_GROUPS}

    no_family = [x["study"] for x in per_study if not x["families"]]

    # ---- results ---------------------------------------------------------
    secondary = read_csv(SEC) if SEC.exists() else []
    prim = {r["analysis_id"]: r for r in read_csv(PRIMARY)}

    def fmt(r, keys=("k", "estimate", "ci_low", "ci_high", "p_value", "tau2", "i2")):
        o = {}
        for k in keys:
            v = r.get(k, "")
            try:
                o[k] = float(v)
            except (TypeError, ValueError):
                o[k] = None
        return o

    payload = {
        "generated_by": "scripts/build_v33_dashboard_data.py",
        "master": V33.name,
        "canonical_studies": len(studies),
        "outcome_rows": len(rows),
        "strict_primary_k": int(float(prim["OP24_PRIM_COMB"]["k"])) if "OP24_PRIM_COMB" in prim else None,
        "result_rob2_coverage": read_csv(ROOT / '09_V34_INTAKE/result_rob2_coverage.csv'),

        "contribution_map": {
            "groups": [{"id": g, "label": l, "definition": d,
                        "n_studies": group_counts[g]} for g, l, d in FAMILY_GROUPS],
            "per_study": per_study,
            "studies_with_no_family": no_family,
            "graph_only_or_unreported": sorted(graph_only),
            "note": ("A study may contribute to several outcome families. Being included in the "
                     "review is not the same as contributing to the primary meta-analysis: most "
                     "included trials report no analysable cumulative 0–24 h opioid dose."),
        },

        "primary": {
            aid: dict(analysis_id=aid, **fmt(r))
            for aid, r in prim.items()
        },

        "secondary": [
            dict(analysis_id=r["analysis_id"], outcome=r["outcome"],
                 figure=('secondary/' + FIGURES[r['analysis_id']]) if r['analysis_id'] in FIGURES else None,
                 measure=r["measure"], model=r["model"], **fmt(r))
            for r in secondary
        ],

        "not_pooled": read_csv(NOTPOOLED) if NOTPOOLED.exists() else [],

        # Interpretive caveats attached to specific analyses. The numbers in
        # each come from the Stata log for that model; the text says what a
        # reader would otherwise have to reconstruct from the weights column.
        "caveats": [
            {
                "analysis_id": "V33_RESCUE_OPIOID_RR_24H",
                "level": "high",
                "text": ("Fragile. Yu 2020 carries 80.8% of the weight in this k=3 model, and it is "
                         "the trial this analysis was newly built around. Leave-one-out: omitting "
                         "Tu 2024 gives RR 0.520 (0.119–2.268, p=0.112) and omitting Yu 2020 gives "
                         "RR 0.433 (0.063–2.995, p=0.114) — the result does not survive either. "
                         "I²=0% reflects three similar point estimates, not a well-estimated "
                         "between-study variance. Hypothesis-generating only."),
            },
            {
                "analysis_id": "V33_GI_DEFECATION_MD",
                "level": "moderate",
                "text": ("Ng 2013 contributes only its sham-controlled contrast; the alternative "
                         "no-acupuncture comparison shares the same EA participants and is not "
                         "counted again. Removing the duplicate materially changes the estimate "
                         "and estimated heterogeneity; it does not justify a certainty upgrade. Result-specific "
                         "RoB 2 is pending for several bowel-function results; flatus or ileus "
                         "assessments do not substitute for defecation assessments."),
            },
            {
                "analysis_id": "V33_INTRAOP_REMI_MD",
                "level": "moderate",
                "text": ("Intraoperative requirement is a different estimand from postoperative "
                         "consumption and says nothing about opioid sparing after surgery. Doses "
                         "are titrated intraoperatively by the anaesthetist, so this outcome is "
                         "vulnerable to performance bias wherever blinding was imperfect. "
                         "Wu 2025 is excluded because its intraoperative doses predate PACU randomization."),
            },
            {
                "analysis_id": "V33_INTRAOP_SUF_MD",
                "level": "moderate",
                "text": "Wu 2025 is excluded: its intraoperative doses were measured before PACU randomization and intervention. These baseline covariates cannot estimate a treatment effect. Intraoperative and postoperative consumption remain separate estimands.",
            },
            {
                "analysis_id": "V33_QOR40_24H_MD",
                "level": "moderate",
                "text": ("k=3 with I²=82.5%. The three trials span a 9-point range of effect "
                         "(+2.6 to +11.7 QoR-40 points) and the interval crosses zero."),
            },
        ],
    }

    header = f"""// v33 DASHBOARD DATA LAYER — generated file, do not hand-edit.
// Regenerate with:  python3 scripts/build_v33_dashboard_data.py
//
// Master : {payload['master']}
// Studies: {payload['canonical_studies']}   Outcome rows: {payload['outcome_rows']}
// Strict primary opioid k: {payload['strict_primary_k']}
//
// Every number is read from the v33 workbook or from a Stata result file.
// None is copied from previous HTML or JS.
window.V33_DATA = """

    OUT.write_text(header + json.dumps(payload, indent=2, ensure_ascii=False) + ";\n",
                   encoding="utf-8")

    print(f"wrote {OUT.relative_to(ROOT)}")
    print(f"  canonical studies ... {payload['canonical_studies']}")
    print(f"  outcome rows ........ {payload['outcome_rows']}")
    print(f"  strict primary k .... {payload['strict_primary_k']}")
    print(f"  secondary analyses .. {len(payload['secondary'])}")
    print(f"  not-pooled register . {len(payload['not_pooled'])}")
    print("\n  contribution map:")
    for g, l, _ in FAMILY_GROUPS:
        print(f"    {l:<34} {group_counts[g]:>3} studies")
    if no_family:
        print(f"    (!) studies with no mapped family: {no_family}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
